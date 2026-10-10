"""Bounded, read-only dataset adapters exposed through the plugin SDK.

Parsing and explicit normalization are shared by preview, comparison and checks.
This module never modifies sources, executes SQL/macros or evaluates expressions.
"""
from __future__ import annotations

import copy
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import json
import math
from pathlib import Path
from typing import Any

from testbox.sdk import PluginError


OPTION_KEYS = {
    "format", "encoding", "delimiter", "quotechar", "escapechar", "record_separator",
    "has_header", "header_row", "start_row", "sheet", "json_path", "columns",
    "skip_empty_rows", "formula_mode", "max_rows", "max_columns", "max_bytes",
    "max_cells", "max_cell_length", "widths", "width_unit",
}
FORMATS = {"auto", "csv", "tsv", "txt", "json", "jsonl", "ndjson", "xlsx", "xlsm", "sql", "fixed"}


def _error(message: str, code: str = "INPUT_INVALID") -> PluginError:
    return PluginError(code, message)


def _integer(options: dict, key: str, default: int, low: int, high: int) -> int:
    value = options.get(key, default)
    if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
        raise _error(f"{key} 必须为 {low} 至 {high} 的整数", "INVALID_PARAMS")
    return value


def _escaped(value: Any, key: str) -> str:
    if not isinstance(value, str):
        raise _error(f"{key} 必须为字符串", "INVALID_PARAMS")
    return value.replace("\\r", "\r").replace("\\n", "\n").replace("\\t", "\t")


def _finite_tree(value: Any, depth: int = 0) -> None:
    if depth > 64:
        raise _error("JSON 嵌套超过 64 层限制", "INPUT_LIMIT_EXCEEDED")
    if isinstance(value, float) and not math.isfinite(value):
        raise _error("输入包含非有限数值")
    if isinstance(value, dict):
        for item in value.values():
            _finite_tree(item, depth + 1)
    elif isinstance(value, list):
        for item in value:
            _finite_tree(item, depth + 1)


def _pairs(pairs: list[tuple[str, Any]]) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            raise _error("JSON 包含重复字段名")
        result[key] = value
    return result


def _json(text: str) -> Any:
    def invalid(_value: str):
        raise _error("JSON 包含非有限数值")
    try:
        value = json.loads(text, object_pairs_hook=_pairs, parse_constant=invalid)
        _finite_tree(value)
        return value
    except (ValueError, RecursionError) as error:
        raise _error("JSON 格式无效或嵌套过深") from error


def _json_path(value: Any, path: Any) -> Any:
    if not isinstance(path, str):
        raise _error("json_path 必须是点分字段路径", "INVALID_PARAMS")
    if path in ("", "$"):
        return value
    parts = path.removeprefix("$.").split(".")
    try:
        for part in parts:
            if isinstance(value, dict):
                value = value[part]
            elif isinstance(value, list) and part.isdigit():
                value = value[int(part)]
            else:
                raise KeyError(part)
    except (KeyError, IndexError, ValueError, OverflowError):
        raise _error("json_path 指定的数据区域不存在") from None
    return value


def _delimited(text: str, delimiter: str, quote: str, escape: str,
               separator: str, max_records: int) -> list[tuple[list[str], int]]:
    """Quote-aware scanner supporting literal multi-character TXT separators."""
    if not delimiter or any(c in delimiter for c in "\r\n"):
        raise _error("delimiter 不得为空或包含换行", "INVALID_PARAMS")
    if len(quote) > 1 or len(escape) > 1 or (quote and quote in delimiter):
        raise _error("quotechar/escapechar 必须为空或单字符且引号不能包含在分隔符中", "INVALID_PARAMS")
    if separator != "auto" and (not separator or separator == delimiter or
                                delimiter.startswith(separator) or separator.startswith(delimiter)):
        raise _error("record_separator 不得为空或与列分隔符重叠", "INVALID_PARAMS")
    records: list[tuple[list[str], int]] = []
    cells: list[str] = []
    field: list[str] = []
    quoted = False
    closed = False
    index = 0
    line = record_line = 1

    def finish_record():
        nonlocal cells, field, closed, record_line
        cells.append("".join(field))
        records.append((cells, record_line))
        if len(records) > max_records:
            raise _error("输入记录数超过限制", "INPUT_LIMIT_EXCEEDED")
        cells, field, closed = [], [], False
        record_line = line

    while index < len(text):
        char = text[index]
        if quoted:
            if escape and char == escape:
                if index + 1 >= len(text):
                    raise _error(f"第 {line} 行转义字符不完整")
                field.append(text[index + 1])
                if text[index + 1] == "\n" or (text[index + 1] == "\r" and not text.startswith("\r\n", index + 1)):
                    line += 1
                index += 2
                continue
            if quote and char == quote:
                if text.startswith(quote + quote, index):
                    field.append(quote)
                    index += 2
                    continue
                quoted, closed = False, True
                index += 1
                continue
            field.append(char)
            if char == "\n" or (char == "\r" and not text.startswith("\r\n", index)):
                line += 1
            index += 1
            continue
        record_sep = ("\r\n" if text.startswith("\r\n", index) else char if char in "\r\n" else "") if separator == "auto" else (separator if text.startswith(separator, index) else "")
        if record_sep:
            line += record_sep.count("\n") or record_sep.count("\r")
            finish_record()
            index += len(record_sep)
            continue
        if text.startswith(delimiter, index):
            cells.append("".join(field))
            field, closed = [], False
            index += len(delimiter)
            continue
        if closed:
            raise _error(f"第 {line} 行引号闭合后存在非法字符")
        if quote and char == quote:
            if field:
                raise _error(f"第 {line} 行未转义引号")
            quoted = True
            index += 1
            continue
        if escape and char == escape:
            if index + 1 >= len(text):
                raise _error(f"第 {line} 行转义字符不完整")
            field.append(text[index + 1])
            if text[index + 1] == "\n" or (text[index + 1] == "\r" and not text.startswith("\r\n", index + 1)):
                line += 1
            index += 2
            continue
        field.append(char)
        if char == "\n" or (char == "\r" and not text.startswith("\r\n", index)):
            line += 1
        index += 1
    if quoted:
        raise _error("输入存在未闭合引号")
    if cells or field or closed:
        finish_record()
    return records


def read_dataset(path: str | Path, options: dict | None = None) -> dict:
    """Return complete bounded records with stable source locations, or fail."""
    if options is None:
        options = {}
    if not isinstance(options, dict) or set(options) - OPTION_KEYS:
        raise _error("解析配置包含不支持的参数", "INVALID_PARAMS")
    maximum = _integer(options, "max_rows", 100_000, 1, 1_000_000)
    max_columns = _integer(options, "max_columns", 256, 1, 4096)
    max_bytes = _integer(options, "max_bytes", 20 * 1024 * 1024, 1, 100 * 1024 * 1024)
    max_cells = _integer(options, "max_cells", 2_000_000, 1, 10_000_000)
    max_length = _integer(options, "max_cell_length", 100_000, 1, 1_000_000)
    header_row = _integer(options, "header_row", 1, 1, 1_000_000)
    start_row = _integer(options, "start_row", 0, 0, 1_000_000)
    for key in ("has_header", "skip_empty_rows"):
        if key in options and not isinstance(options[key], bool):
            raise _error(f"{key} 必须为布尔值", "INVALID_PARAMS")
    source = Path(path)
    if not source.is_file():
        raise _error("输入文件不存在", "INPUT_NOT_FOUND")
    if source.stat().st_size > max_bytes:
        raise _error("输入文件大小超过限制", "INPUT_LIMIT_EXCEEDED")
    fmt = options.get("format", "auto")
    if not isinstance(fmt, str) or fmt not in FORMATS:
        raise _error("不支持的输入格式", "INVALID_PARAMS")
    if fmt == "auto":
        fmt = source.suffix.lower().lstrip(".")
        if fmt == "ddl":
            fmt = "sql"
        if fmt == "xls":
            raise _error("旧版 .xls 尚不支持，请转换为 .xlsx")
        if fmt not in FORMATS - {"auto"}:
            raise _error("无法根据扩展名确定输入格式，请明确设置 format")
    rows: list[dict] = []
    locations: list[dict] = []
    columns: list[str] = []
    warnings: list[str] = []
    sheet_name = None
    cells_count = 0

    def add_row(row: dict, location: dict):
        nonlocal cells_count
        if not isinstance(row, dict) or not all(isinstance(key, str) and key for key in row):
            raise _error("每条 JSON 数据必须是字段名非空的对象")
        _finite_tree(row)
        for key, value in row.items():
            if key not in columns:
                columns.append(key)
            if len(key) > max_length or (len(value) if isinstance(value, str) else len(json.dumps(value, ensure_ascii=False))) > max_length:
                raise _error("单元格或字段名长度超过限制", "INPUT_LIMIT_EXCEEDED")
        cells_count += len(row)
        if len(columns) > max_columns or cells_count > max_cells or len(rows) >= maximum:
            raise _error("输入行数、列数或单元格数超过限制", "INPUT_LIMIT_EXCEEDED")
        rows.append(row)
        locations.append({"source": str(source), **location})

    if fmt in {"json", "jsonl", "ndjson", "sql", "csv", "tsv", "txt", "fixed"}:
        encoding = options.get("encoding", "utf-8-sig")
        if not isinstance(encoding, str):
            raise _error("encoding 必须为字符串", "INVALID_PARAMS")
        try:
            with source.open("rb") as stream:
                payload = stream.read(max_bytes + 1)
            if len(payload) > max_bytes:
                raise _error("输入文件大小超过限制", "INPUT_LIMIT_EXCEEDED")
            text = payload.decode(encoding)
        except (UnicodeError, LookupError):
            raise _error("输入编码无效或与文件不匹配，请调整 encoding") from None
        if fmt == "fixed":
            widths = options.get("widths")
            unit = options.get("width_unit", "characters")
            if (not isinstance(widths, list) or not widths or len(widths) > max_columns
                    or any(type(w) is not int or not 1 <= w <= max_length for w in widths)
                    or sum(widths) > max_length * max_columns):
                raise _error("widths 必须为有界正整数宽度列表", "INVALID_PARAMS")
            if not isinstance(unit, str) or unit not in {"characters", "bytes"}:
                raise _error("width_unit 必须为 characters 或 bytes", "INVALID_PARAMS")
            if unit == "bytes" and encoding.lower().replace("_", "-") not in {"utf-8", "utf-8-sig", "ascii", "gbk", "gb2312", "gb18030", "big5", "latin-1", "latin1", "cp1252"}:
                raise _error("字节定宽仅支持 UTF-8、GBK/GB18030、Big5 或单字节编码", "INVALID_PARAMS")
            codec = encoding.lower().replace("_", "-").replace("-sig", "")
            separator = _escaped(options.get("record_separator", "auto"), "record_separator")
            if not separator:
                raise _error("record_separator 不得为空", "INVALID_PARAMS")
            # Split only physical CR/LF records by default; other characters are data.
            import re
            raw = payload.removeprefix(b"\xef\xbb\xbf") if encoding.lower().replace("_", "-") == "utf-8-sig" else payload
            value = raw if unit == "bytes" else text
            if separator == "auto":
                records_raw = re.split(b"\r\n|\r|\n" if unit == "bytes" else r"\r\n|\r|\n", value)
            else:
                records_raw = value.split(separator.encode(codec) if unit == "bytes" else separator)
            if records_raw and not records_raw[-1]:
                records_raw.pop()  # Final terminator is not an extra empty record.
            records = []
            line = 1
            for number, record in enumerate(records_raw, 1):
                location = line
                line += record.count(b"\n" if unit == "bytes" else "\n") + (1 if separator == "auto" else separator.count("\n"))
                if not record and options.get("skip_empty_rows", True):
                    continue
                if len(record) != sum(widths):
                    raise _error(f"第 {number} 条记录宽度为 {len(record)}，预期 {sum(widths)}；拒绝截断或补齐")
                cells, offset = [], 0
                for width in widths:
                    cell = record[offset:offset + width]
                    offset += width
                    if unit == "bytes":
                        try:
                            cell = cell.decode(codec)
                        except UnicodeError:
                            raise _error(f"第 {number} 条记录字段边界切断编码字符") from None
                    cells.append(cell)
                records.append((cells, location))
                if len(records) > maximum + header_row + start_row:
                    raise _error("输入记录数超过限制", "INPUT_LIMIT_EXCEEDED")
            field_ranges = None
            def add_fixed(row, location):
                nonlocal field_ranges
                if field_ranges is None:
                    field_ranges, offset = {}, 1
                    for name, width in zip(columns, widths):
                        field_ranges[name] = {"start": offset, "width": width}
                        offset += width
                add_row(row, {**location, "width_unit": unit, "field_ranges": field_ranges})
            _read_matrix(records, options, header_row, start_row, columns, add_fixed, max_columns)
        elif fmt == "sql":
            # SQL text is not a spreadsheet cell; allow bounded full source here.
            columns, rows = ["sql"], [{"sql": text}]
            locations = [{"source": str(source), "row": 1}]
        elif fmt == "json":
            data = _json_path(_json(text), options.get("json_path", ""))
            if isinstance(data, dict):
                data = [data]
            if not isinstance(data, list):
                raise _error("JSON 数据区域必须是对象或对象数组")
            for index, row in enumerate(data):
                add_row(row, {"row": index + 1, "path": f"{options.get('json_path') or '$'}[{index}]"})
        elif fmt in {"jsonl", "ndjson"}:
            for index, line in enumerate(text.splitlines()):
                if not line.strip():
                    if options.get("skip_empty_rows", True):
                        continue
                    raise _error(f"JSONL 第 {index + 1} 行为空")
                row = _json_path(_json(line), options.get("json_path", ""))
                add_row(row, {"row": index + 1})
        else:
            if "delimiter" in options:
                delimiter = _escaped(options["delimiter"], "delimiter")
            elif fmt == "txt":
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",\t|;").delimiter
                    warnings.append("TXT 分隔符由启发式识别，请在预览中确认或显式设置 delimiter")
                except csv.Error:
                    raise _error("无法确定 TXT 分隔符，请明确设置 delimiter") from None
            else:
                delimiter = "\t" if fmt == "tsv" else ","
            if fmt != "txt" and len(delimiter) != 1:
                raise _error("CSV/TSV 的 delimiter 必须是单字符；多字符请使用 txt 模式", "INVALID_PARAMS")
            separator = _escaped(options.get("record_separator", "auto"), "record_separator")
            records = _delimited(text, delimiter, _escaped(options.get("quotechar", '"'), "quotechar"),
                                 _escaped(options.get("escapechar", ""), "escapechar"), separator,
                                 min(maximum + header_row + start_row + 1, 1_000_002))
            records = [(cells, line) for cells, line in records if not options.get("skip_empty_rows", True) or cells != [""]]
            _read_matrix(records, options, header_row, start_row, columns, add_row, max_columns)
    else:
        if not isinstance(options.get("formula_mode", "formula"), str) or options.get("formula_mode", "formula") not in {"formula", "cached"}:
            raise _error("formula_mode 必须为 formula 或 cached", "INVALID_PARAMS")
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            raise _error("Excel 读取需要已有可选依赖 openpyxl（evidence extra）", "DEPENDENCY_MISSING") from None
        # Enforce expansion bounds before handing the ZIP to an XML parser.
        import zipfile
        try:
            with zipfile.ZipFile(source) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise _error("Excel 解压大小超过限制", "INPUT_LIMIT_EXCEEDED")
            workbook = load_workbook(source, read_only=True,
                                     data_only=options.get("formula_mode", "formula") == "cached")
            try:
                selection = options.get("sheet", "")
                if not isinstance(selection, str):
                    raise _error("sheet 必须为工作表名称", "INVALID_PARAMS")
                if selection and selection not in workbook.sheetnames:
                    raise _error("指定工作表不存在")
                sheet = workbook[selection] if selection else workbook.active
                sheet_name = sheet.title
                if not selection and len(workbook.sheetnames) > 1:
                    warnings.append("仅读取活动工作表；多工作表请明确指定 sheet")
                if options.get("formula_mode") == "cached":
                    warnings.append("使用公式缓存值，不重新计算；无缓存与空值无法仅凭此模式区分")
                if (sheet.max_row or 0) > maximum + header_row + start_row or (sheet.max_column or 0) > max_columns:
                    raise _error("Excel 行列数超过限制", "INPUT_LIMIT_EXCEEDED")
                records = []
                for index, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                    converted = [_excel_value(value) for value in cells]
                    if options.get("skip_empty_rows", True) and all(value is None for value in converted):
                        continue
                    records.append((converted, index))
                    if len(records) > maximum + header_row + start_row:
                        raise _error("Excel 行数超过限制", "INPUT_LIMIT_EXCEEDED")
                _read_matrix(records, options, header_row, start_row, columns,
                             lambda row, loc: add_row(row, {"sheet": sheet.title, **loc}), max_columns,
                             excel=True)
            finally:
                workbook.close()
        except PluginError:
            raise
        except Exception as error:
            raise _error(f"Excel 读取失败（{type(error).__name__}）") from None
    if not columns and options.get("columns"):
        columns = _headers(options["columns"], max_columns)
    if any(len(column) > max_length for column in columns):
        raise _error("字段名长度超过限制", "INPUT_LIMIT_EXCEEDED")
    return {"columns": columns, "rows": rows, "locations": locations, "format": fmt,
            "warnings": warnings, "source": str(source), "complete": True, "sheet": sheet_name}


def _excel_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _headers(values: Any, limit: int) -> list[str]:
    if not isinstance(values, list) or not all(isinstance(value, str) and value for value in values):
        raise _error("表头/columns 必须是非空字符串列表")
    if len(values) > limit or len(values) != len(set(values)):
        raise _error("表头列数超限或存在重复字段名")
    return list(values)


def _read_matrix(records: list, options: dict, header_row: int, start_row: int,
                 columns: list, add_row, max_columns: int, excel: bool = False):
    has_header = options.get("has_header", True)
    if has_header:
        # Excel positions are physical rows; delimited text positions are records.
        selected = next((item for item in records if item[1] == header_row), None) if excel else (records[header_row - 1] if len(records) >= header_row else None)
        if selected is None:
            if not records:
                raise _error("文件为空或没有表头；无表头文件请配置 has_header=false 和 columns")
            raise _error("header_row 超出输入范围")
        columns.extend(_headers([str(value) if value is not None else "" for value in selected[0]], max_columns))
        offset = (records.index(selected) + 1) if not start_row else (next((i for i, item in enumerate(records) if item[1] >= start_row), len(records)) if excel else start_row - 1)
        if start_row and (start_row <= header_row):
            raise _error("start_row 必须位于表头之后", "INVALID_PARAMS")
    else:
        supplied = options.get("columns")
        if supplied is None:
            if not records:
                return
            supplied = [f"column_{i + 1}" for i in range(len(records[0][0]))]
        columns.extend(_headers(supplied, max_columns))
        offset = next((i for i, item in enumerate(records) if item[1] >= start_row), len(records)) if excel and start_row else max(0, start_row - 1)
    for record, (cells, line) in enumerate(records[offset:], offset + 1):
        if len(cells) != len(columns):
            raise _error(f"第 {line} 行字段数量与表头不一致")
        add_row(dict(zip(columns, cells)), {"row": line, **({} if excel else {"record": record})})


def normalize_dataset(dataset: dict, rules: dict | None = None) -> dict:
    """Apply only explicit, deterministic rules to an independent dataset copy."""
    if rules is None:
        rules = {}
    allowed = {"column_mapping", "trim", "casefold", "null_values", "types"}
    if not isinstance(rules, dict) or set(rules) - allowed:
        raise _error("标准化配置包含不支持的参数", "INVALID_PARAMS")
    for key in ("trim", "casefold"):
        if key in rules and not isinstance(rules[key], bool):
            raise _error(f"{key} 必须为布尔值", "INVALID_PARAMS")
    mapping, types = rules.get("column_mapping", {}), rules.get("types", {})
    if not isinstance(mapping, dict) or not all(isinstance(key, str) and isinstance(value, str) and value for key, value in mapping.items()):
        raise _error("column_mapping 必须是字段名映射", "INVALID_PARAMS")
    if not isinstance(types, dict) or not all(isinstance(key, str) and isinstance(value, str) and value in {"string", "integer", "decimal", "boolean", "date", "datetime"} for key, value in types.items()):
        raise _error("types 包含不支持的类型", "INVALID_PARAMS")
    nulls = rules.get("null_values", [])
    if not isinstance(nulls, list) or not all(isinstance(item, str) for item in nulls):
        raise _error("null_values 必须是字符串列表", "INVALID_PARAMS")
    if set(mapping) - set(dataset["columns"]):
        raise _error("column_mapping 引用了不存在的输入字段", "INVALID_PARAMS")
    names = [mapping.get(column, column) for column in dataset["columns"]]
    if len(names) != len(set(names)):
        raise _error("标准化后字段名冲突", "INVALID_PARAMS")
    if set(types) - set(names):
        raise _error("types 引用了不存在的标准化字段", "INVALID_PARAMS")
    result = copy.deepcopy(dataset)
    result["columns"], result["types"] = names, dict(types)
    result["normalization"] = copy.deepcopy(rules)
    result["rows"] = []
    for index, row in enumerate(dataset["rows"]):
        normalized = {}
        for original, value in row.items():
            value = copy.deepcopy(value)
            key = mapping.get(original, original)
            if isinstance(value, str):
                if rules.get("trim"):
                    value = value.strip()
                if rules.get("casefold"):
                    value = value.casefold()
                if value in nulls:
                    value = None
            if value is not None and key in types:
                try:
                    value = _convert(value, types[key])
                except (ValueError, InvalidOperation, TypeError, OverflowError):
                    location = dataset["locations"][index]
                    raise _error(f"第 {location.get('row', index + 1)} 行字段 {key} 类型转换失败") from None
            normalized[key] = value
        result["rows"].append(normalized)
    return result


def _convert(value: Any, kind: str) -> Any:
    if kind == "string":
        if isinstance(value, (dict, list)):
            raise ValueError()
        return str(value)
    if kind in {"integer", "decimal"}:
        if isinstance(value, (bool, dict, list)):
            raise ValueError()
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValueError()
        if abs(number.adjusted()) > 1000:
            raise ValueError()
        if kind == "integer":
            if number != number.to_integral_value():
                raise ValueError()
            return int(number)
        if abs(number.adjusted()) > 1000:
            raise ValueError()
        # format is exact and does not round under Decimal's ambient context.
        formatted = format(number, "f")
        if "." in formatted:
            formatted = formatted.rstrip("0").rstrip(".")
        return "0" if number == 0 else formatted
    if kind == "boolean":
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value in {"true", "false"}:
            return value == "true"
        raise ValueError()
    if not isinstance(value, str):
        raise ValueError()
    return (date.fromisoformat(value) if kind == "date" else datetime.fromisoformat(value)).isoformat()
