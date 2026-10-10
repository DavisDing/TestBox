"""Bounded sequential orchestration; original readers and business rules stay authoritative."""
from __future__ import annotations

from collections import Counter
from dataclasses import replace
import csv
import io
import json
from pathlib import Path
import zipfile

from testbox.sdk import PluginError, Result, Task

MAX_ITEMS = 100


def invalid(message):
    raise PluginError("INVALID_PARAMS", message)


def sheet_names(path, options):
    """Inspect read-only workbook metadata with the same input/ZIP bounds as the reader."""
    from testbox.tabular import _integer
    maximum = _integer(options, "max_bytes", 20 * 1024 * 1024, 1, 100 * 1024 * 1024)
    source = Path(path)
    if not source.is_file():
        raise PluginError("INPUT_NOT_FOUND", "工作簿不存在")
    if source.stat().st_size > maximum:
        raise PluginError("INPUT_LIMIT_EXCEEDED", "工作簿大小超过限制")
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        raise PluginError("DEPENDENCY_MISSING", "Excel 读取需要 openpyxl") from None
    try:
        with zipfile.ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise PluginError("INPUT_LIMIT_EXCEEDED", "Excel 解压大小超过限制")
        workbook = load_workbook(source, read_only=True)
        try:
            names = workbook.sheetnames
            if len(names) > MAX_ITEMS:
                raise PluginError("INPUT_LIMIT_EXCEEDED", "工作表数量超过100项限制")
            return names
        finally:
            workbook.close()
    except PluginError:
        raise
    except Exception as error:
        raise PluginError("INPUT_INVALID", f"工作簿读取失败（{type(error).__name__}）") from None


def selection(params, side):
    key = "inputs" if side == "input" else side + "_inputs"
    paths = params.get(key, [])
    if not isinstance(paths, list) or len(paths) > MAX_ITEMS or any(not isinstance(p, str) or not p for p in paths):
        invalid(f"{key} 必须为最多100个非空文件路径")
    single = params.get(side)
    if single and paths:
        invalid(f"{side} 与 {key} 不能同时提供")
    if single and (not isinstance(single, str)):
        invalid(f"{side} 必须为文件路径")
    if len(set(paths)) != len(paths):
        invalid(f"{key} 存在重复来源文件")
    if paths or single:
        text = "text" if side == "input" else side + "_text"
        if params.get(text):
            invalid("批量文件与直接 SQL 文本不能混用")
    return paths or ([single] if single else [])


def expand(paths, options):
    if not isinstance(options, dict):
        invalid("解析配置必须为对象")
    options = dict(options)
    has_sheets = "sheets" in options
    sheets = options.pop("sheets", None)
    if has_sheets and sheets is None:
        invalid('sheets 不能为 null，请使用 "all" 或名称数组')
    if sheets is not None:
        if options.get("sheet"):
            invalid("sheet 与 sheets 不能同时提供")
        if sheets != "all" and (not isinstance(sheets, list) or not sheets
                or len(sheets) > MAX_ITEMS or any(not isinstance(s, str) or not s for s in sheets)
                or len(set(sheets)) != len(sheets)):
            invalid('sheets 必须为 "all" 或不重复的工作表名称数组')
    units = []
    if len(paths) > MAX_ITEMS:
        raise PluginError("INPUT_LIMIT_EXCEEDED", "输入文件超过100个")
    for path in paths:
        fmt = options.get("format", "auto")
        excel = fmt in {"xlsx", "xlsm"} or (fmt == "auto" and Path(path).suffix.lower() in {".xlsx", ".xlsm"})
        names = [options.get("sheet", "")]
        failure = None
        if sheets is not None:
            if not excel:
                invalid("sheets 仅适用于 Excel；混合文件批次请分别执行")
            try:
                available = sheet_names(path, options)
                names = available if sheets == "all" else sheets
            except PluginError as error:
                names, failure = [""], error
        for name in names:
            selected = dict(options)
            if name:
                selected["sheet"] = name
            units.append({"path": path, "name": Path(path).stem, "sheet": name,
                          "options": selected, "failure": failure})
            if len(units) > MAX_ITEMS:
                raise PluginError("INPUT_LIMIT_EXCEEDED", "文件×Sheet展开后超过100项；请缩小批次")
    return units


def run_batch(plugin, command, params):
    compare = command in {"data.compare", "sql.diff"}
    sides = ["left", "right"] if compare else ["input"]
    batch = params.get("batch", {})
    if not isinstance(batch, dict) or set(batch) - {"pairing", "sheet_mapping"}:
        invalid("batch 仅支持 pairing 和 sheet_mapping")
    active = bool(batch) or any(params.get("inputs" if side == "input" else side + "_inputs")
            or isinstance(params.get("options" if side == "input" else side + "_options"), dict)
            and "sheets" in params.get("options" if side == "input" else side + "_options", {}) for side in sides)
    # Do not alter established single-file / direct SQL contracts.
    if not active:
        return None
    pairing = batch.get("pairing", "name")
    mapping = batch.get("sheet_mapping", {})
    if not isinstance(pairing, str) or pairing not in {"name", "position"} or not isinstance(mapping, dict) or any(
            not isinstance(k, str) or not k or not isinstance(v, str) or not v for k, v in mapping.items()):
        invalid("pairing 必须为 name/position；sheet_mapping 为非空Sheet名称映射")
    if len(set(mapping.values())) != len(mapping) or (mapping and pairing != "name"):
        invalid("Sheet映射必须一对一且使用name配对")
    expanded = []
    for side in sides:
        paths = selection(params, side)
        if not paths:
            invalid(f"请选择 {side} 的单文件或批量文件")
        expanded.append(expand(paths, params.get("options" if side == "input" else side + "_options", {})))
    plans = []
    if compare:
        left, right = expanded
        if pairing == "position":
            for index in range(max(len(left), len(right))):
                plans.append((left[index] if index < len(left) else None, right[index] if index < len(right) else None, None))
        else:
            # A single file pair can have different filenames, as in the old API.
            single_pair = len({u["path"] for u in left}) == len({u["path"] for u in right}) == 1
            def key(unit, left_side=False):
                return ("" if single_pair else unit["name"], mapping.get(unit["sheet"], unit["sheet"]) if left_side else unit["sheet"])
            lc, rc = Counter(key(u, True) for u in left), Counter(key(u) for u in right)
            used = set()
            for unit in left:
                k = key(unit, True)
                matches = [i for i, u in enumerate(right) if key(u) == k]
                reason = "AMBIGUOUS_PAIR" if lc[k] > 1 or rc[k] > 1 else None
                if reason:
                    used.update(matches)
                    plans.append((unit, None, reason))
                elif matches:
                    used.add(matches[0]); plans.append((unit, right[matches[0]], None))
                else:
                    plans.append((unit, None, "UNMATCHED"))
            for i, unit in enumerate(right):
                if i not in used:
                    plans.append((None, unit, "AMBIGUOUS_PAIR" if rc[key(unit)] > 1 else "UNMATCHED"))
    else:
        plans = [(unit, None, None) for unit in expanded[0]]
    if len(plans) > MAX_ITEMS:
        raise PluginError("INPUT_LIMIT_EXCEEDED", "批次配对结果超过100项；请缩小批次")
    original = plugin.context
    items, files = [], []
    extras = {"inputs", "left_inputs", "right_inputs", "batch"}
    for index, (first, second, reason) in enumerate(plans, 1):
        item = {"index": index, "left" if compare else "input":
                {k: first[k] for k in ("name", "sheet")} if first else None}
        if compare:
            item["right"] = {k: second[k] for k in ("name", "sheet")} if second else None
        if reason or (compare and (first is None or second is None)):
            item.update(status="unmatched" if reason in {None, "UNMATCHED"} else "failed", code=reason or "UNMATCHED")
            items.append(item); continue
        child = {k: v for k, v in params.items() if k not in extras}
        for side, unit in zip(sides, (first, second)):
            child[side] = unit["path"]
            child["options" if side == "input" else side + "_options"] = unit["options"]
        plugin.context = replace(original, task=Task(f"{original.task.id}-item-{index:03d}"))
        try:
            for unit in (first, second):
                if unit and unit["failure"]:
                    raise unit["failure"]
            result = plugin.execute_one(command, child)
            item.update(status=result.status, message=result.message, files=result.files, data=result.data, warnings=result.warnings)
            files.extend(result.files)
            if command.endswith("preview") and result.status == "success":
                # Store display data separately, not hundreds of tables in the Host response.
                name = f"{plugin.context.task.id}-display.json"
                original.files.write_text(name, json.dumps(result.data, ensure_ascii=False, allow_nan=False))
                item["preview_file"] = name; files.append(name)
        except PluginError as error:
            item.update(status="failed", code=error.code, message=str(error))
        except Exception as error:
            item.update(status="failed", code="INPUT_INVALID", message=f"单项执行失败（{type(error).__name__}）")
        finally:
            plugin.context = original
        items.append(item)
    failed = sum(i["status"] != "success" for i in items)
    for item in items:
        # Limit repeated per-item metadata in summaries. Full data remains in child reports.
        if item.get("message"):
            item["message"] = item["message"][:2000]
        if "data" in item:
            warnings_count = item["data"].get("warnings")
            item["data"] = {k: v for k, v in item["data"].items() if k not in {"rows", "raw_rows", "locations", "structure", "unknown", "differences", "warnings", "output_files"}}
            if type(warnings_count) is int:
                item["data"]["warnings"] = warnings_count
        if "warnings" in item:
            item["warnings"] = item["warnings"][:100]
    summary = {"batch": True, "total": len(items), "succeeded": len(items) - failed, "failed": failed,
               "execution_complete": failed == 0, "items": items}
    if command == "data.compare":
        summary["different"] = sum(i.get("data", {}).get("equal") is False for i in items)
        summary["equal"] = not failed and all(i.get("data", {}).get("equal") is True for i in items)
    elif command == "data.check":
        summary["quality_failed"] = sum(i.get("data", {}).get("passed") is False for i in items)
        summary["passed"] = not failed and all(i.get("data", {}).get("passed") is True for i in items)
    elif command == "sql.diff":
        verdicts = [i.get("data", {}).get("verdict") for i in items]
        summary["verdict"] = "inconclusive" if failed else "different" if "different" in verdicts else "equal" if all(v == "equal" for v in verdicts) else "inconclusive"
    name = f"{original.task.id}-batch.json"
    original.files.write_text(name, json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False)); files.append(name)
    csv_output = io.StringIO()
    writer = csv.writer(csv_output, lineterminator="\n")
    writer.writerow(["index", "left_or_input", "left_sheet", "right", "right_sheet", "status", "equal", "passed", "verdict", "code", "message"])
    def safe_cell(value):
        text = str(value if value is not None else "")
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")) else text
    for item in items:
        left = item.get("left") or item.get("input") or {}
        right = item.get("right") or {}
        data = item.get("data", {})
        writer.writerow([safe_cell(v) for v in [item["index"], left.get("name"), left.get("sheet"), right.get("name"), right.get("sheet"), item["status"], data.get("equal"), data.get("passed"), data.get("verdict"), item.get("code"), item.get("message")]])
    csv_name = f"{original.task.id}-batch.csv"
    original.files.write_text(csv_name, csv_output.getvalue(), encoding="utf-8-sig"); files.append(csv_name)
    # Transport only bounded per-item summaries; full reports remain registered files.
    display = dict(summary)
    display["items"] = [{k: v for k, v in i.items() if k not in {"data", "warnings"}} |
                        {"outcome": {k: i.get("data", {}).get(k) for k in ("equal", "passed", "verdict", "complete", "difference_count", "errors", "warnings") if k in i.get("data", {})}}
                        for i in items]
    return Result("failed" if failed else "success", f"批量完成：{len(items)-failed}/{len(items)} 项执行成功，{failed} 项失败或未配对",
                  display, files, ["部分项目失败或未配对，详见批量报告"] if failed else [])
