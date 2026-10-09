"""Verify an installed distribution from an empty cwd, not the source checkout.

Usage: python scripts/smoke_installed.py --python /path/to/clean/venv/python
Add --gui / --evidence only when the target environment has those extras.
All task data, fixture files and exports are created under TemporaryDirectory.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import zipfile


COMMANDS = {"data.mock", "sql.parse", "sql.select", "evidence.build", "data.preview", "data.compare", "data.check", "sql.diff", "sql.preview"}


def verify_installed(python: Path, *, gui: bool = False, evidence: bool = False) -> dict:
    # Do not resolve the venv executable symlink: that would run global Python.
    python = python.absolute()
    if not python.is_file():
        raise ValueError(f"Target Python does not exist: {python}")
    source = Path(__file__).resolve().parents[1]
    with tempfile.TemporaryDirectory(prefix="testbox-installed-smoke-") as temporary:
        root = Path(temporary).resolve()
        cwd = root / "empty-cwd"
        cwd.mkdir()
        environment = os.environ.copy()
        for name in ("PYTHONPATH", "PYTHONHOME"):
            environment.pop(name, None)
        environment.update({
            "PYTHONNOUSERSITE": "1", "PYTHONUTF8": "1", "QT_QPA_PLATFORM": "offscreen",
            "HOME": str(root / "home"), "USERPROFILE": str(root / "home"),
            "LOCALAPPDATA": str(root / "data"), "XDG_DATA_HOME": str(root / "data"),
        })
        (root / "home").mkdir()

        def call(arguments: list[str]) -> subprocess.CompletedProcess:
            process = subprocess.run(
                [str(python), "-I", *arguments], cwd=cwd, env=environment,
                capture_output=True, text=True, encoding="utf-8", timeout=60,
            )
            if process.returncode:
                raise RuntimeError(f"Installed smoke failed ({process.returncode}): {arguments}\n{process.stdout}\n{process.stderr}")
            return process

        def code(script: str):
            return json.loads(call(["-c", script]).stdout)

        def cli(*arguments: str):
            return json.loads(call(["-m", "testbox.cli", "--json", *arguments]).stdout)

        initial = code('''import hashlib, json
from pathlib import Path
import testbox
from testbox.core.runtime import Runtime
r = Runtime()
try:
    bundled = r.bundled_plugins_dir
    hashes = {str(p.relative_to(bundled)): hashlib.sha256(p.read_bytes()).hexdigest()
              for p in bundled.rglob('*') if p.is_file() and '__pycache__' not in p.parts}
    print(json.dumps({'module': testbox.__file__, 'diagnostics': r.get_runtime_diagnostics(), 'hashes': hashes}))
finally: r.close()
''')
        module = Path(initial["module"]).resolve()
        if source == module or source in module.parents:
            raise AssertionError(f"Smoke accidentally imported source checkout: {module}")
        bundled = Path(initial["diagnostics"]["bundled_plugins_dir"])
        workspace = Path(initial["diagnostics"]["workspace_dir"])
        user_plugins = Path(initial["diagnostics"]["plugins_dir"])
        if root not in workspace.parents or root not in user_plugins.parents:
            raise AssertionError("Installed task/user-plugin storage escaped temporary user data")
        if bundled.name != "_bundled_plugins" or module.parent not in bundled.parents:
            raise AssertionError("Official plugins were not loaded from installed package resources")
        for name in ("logo.ico", "logo.svg"):
            if not (module.parent / "assets" / name).is_file():
                raise AssertionError(f"Missing installed GUI asset: {name}")
        metadata = json.loads((bundled / "data-generator" / "data" / "metadata.json").read_text())
        divisions = metadata["administrative_divisions"]
        for key in ("mainland", "hk_mo_tw"):
            data_file = bundled / "data-generator" / "data" / divisions[f"{key}_file"]
            if hashlib.sha256(data_file.read_bytes()).hexdigest() != divisions[f"{key}_sha256"]:
                raise AssertionError(f"Bundled third-party resource hash mismatch: {data_file}")

        listing = cli("plugin", "list")
        if {item["command"] for item in listing["commands"]} != COMMANDS or listing["unavailable"]:
            raise AssertionError(f"Incomplete installed command registry: {listing}")
        entry = python.parent / ("testbox.exe" if os.name == "nt" else "testbox")
        entry_process = subprocess.run([str(entry), "--json", "plugin", "list"],
                                       cwd=cwd, env=environment, capture_output=True,
                                       text=True, encoding="utf-8", timeout=30)
        if entry_process.returncode or json.loads(entry_process.stdout) != listing:
            raise AssertionError(f"Installed console entry point failed: {entry_process.stderr}")
        for command in sorted(COMMANDS):
            schema = code(f'''import json
from testbox.core.runtime import Runtime
r=Runtime()
try: print(json.dumps(r.get_command_schema({command!r})))
finally: r.close()
''')
            if schema.get("type") != "object":
                raise AssertionError(f"Missing installed schema: {command}")

        # Use the real built-in customer template to exercise fixed data/config.
        params = {"count": 3, "format": "json", "seed": 17, "template": "retail_customer"}

        def run(command: str, values: dict):
            arguments = ["run", command]
            for name, value in values.items():
                arguments += ["--set", f"{name}={json.dumps(value, ensure_ascii=False)}"]
            result = cli(*arguments)
            if result["status"] != "success":
                raise AssertionError(result)
            if Path(result["workspace"]).parent != workspace:
                raise AssertionError("Task did not use installed user workspace")
            task = Path(result["workspace"])
            for name in ("manifest.json", "result.json", "report.md"):
                if not (task / name).is_file():
                    raise AssertionError(f"Missing trace artifact: {name}")
            return result

        first, second = run("data.mock", params), run("data.mock", params)
        generated = [Path(result["workspace"]) / "output" / result["files"][0] for result in (first, second)]
        if generated[0].read_bytes() != generated[1].read_bytes():
            raise AssertionError("Installed Host data generation is not seed reproducible")
        ddl = cwd / "schema.sql"
        ddl.write_text("CREATE TABLE users (id INT PRIMARY KEY, name VARCHAR(40));", encoding="utf-8")
        parsed = run("sql.parse", {"input": str(ddl), "format": "json"})
        fields = Path(parsed["workspace"]) / "output" / parsed["files"][0]
        selected = run("sql.select", {"input": str(fields), "dialect": "mysql"})
        select_file = Path(selected["workspace"]) / "output" / selected["files"][0]
        if "SELECT" not in select_file.read_text() or "users" not in select_file.read_text():
            raise AssertionError("Installed SQL handoff produced unexpected text")
        stored = cli("task", "result", selected["task_id"])
        if stored["files"] != selected["files"]:
            raise AssertionError("Installed task result could not be read back")
        archive = cwd / "export.zip"
        cli("task", "export", selected["task_id"], "--archive", "--output", str(archive))
        with zipfile.ZipFile(archive) as exported:
            if exported.namelist() != selected["files"]:
                raise AssertionError("Installed export did not preserve declared artifact list")

        # Exercise new business commands through installed CLI -> real Host.
        previewed = run("data.preview", {"input": str(generated[0]), "sample_rows": 1})
        if previewed["data"]["row_count"] != 3 or not previewed["data"]["sample_truncated"]:
            raise AssertionError("Installed preview did not parse the complete generated dataset")
        compared = run("data.compare", {"left": str(generated[0]), "right": str(generated[1]), "mode": "multiset"})
        if not compared["data"]["equal"]:
            raise AssertionError("Installed comparison did not preserve equal repeated data")
        checked = run("data.check", {"input": str(generated[0]), "rules": [{"type": "row_count", "min": 3, "max": 3}]})
        if not checked["data"]["passed"]:
            raise AssertionError("Installed data quality row count failed")
        diffed = run("sql.diff", {"left": str(ddl), "right": str(ddl)})
        if diffed["data"]["verdict"] != "equal":
            raise AssertionError("Installed identical DDL comparison failed")
        sql_previewed = run("sql.preview", {"input": str(ddl)})
        if sql_previewed["data"]["total_rows"] != 2:
            raise AssertionError("Installed SQL preview lost declared columns")

        # A user-installed same-name override can be removed; immutable official
        # resources remain available afterwards and are never rewritten.
        package = cwd / "data-generator.zip"
        cli("plugin", "package", str(bundled / "data-generator"), "--output", str(package))
        cli("plugin", "install", str(package))
        if not (user_plugins / "data-generator" / "manifest.yaml").is_file():
            raise AssertionError("Plugin override did not install into user data")
        cli("plugin", "uninstall", "data-generator")
        if {item["command"] for item in cli("plugin", "list")["commands"]} != COMMANDS:
            raise AssertionError("Removing user override removed an official plugin")
        summary = {"cli": "passed", "resources": "passed", "sql_handoff": "passed",
                   "reproducibility": "passed", "export": "passed", "plugin_override": "passed",
                   "gui": "not_requested", "evidence": "not_requested",
                   "data_preview": "passed", "data_compare": "passed", "data_check": "passed", "sql_diff": "passed"}
        if evidence:
            fixtures = code('''import json
from pathlib import Path
from openpyxl import Workbook
from PIL import Image
cases = Path.cwd() / 'cases.xlsx'
book = Workbook(); sheet = book.active
sheet.append(['Case Name', 'Check Point', 'Step', 'Desc', 'Expected', 'Status'])
sheet.append(['Smoke case', 'TEST DATA ONLY', '1', 'Synthetic step', 'Synthetic result', ''])
book.save(cases); book.close()
shot = Path.cwd() / 'shot.png'; Image.new('RGB', (320, 180), 'white').save(shot)
print(json.dumps({'input': str(cases), 'screenshots': [str(shot)]}))
''')
            evidence_result = run("evidence.build", {**fixtures, "interactive": False, "update_excel": True})
            output = Path(evidence_result["workspace"]) / "output"
            check = code(f'''import json
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
output = Path({str(output)!r})
reports = list((output / 'reports').glob('*.docx'))
assert len(reports) == 1
assert any('TEST DATA ONLY' in p.text for p in Document(reports[0]).paragraphs)
book = load_workbook(output / 'executed-cases.xlsx')
assert book.active['F2'].value == '已执行'
book.close()
print(json.dumps({{'status': 'passed'}}))
''')
            summary["evidence"] = check["status"]
        if gui:
            gui_result = code('''import json, time
from PySide6 import QtCore, QtWidgets
from testbox.gui import MainWindow
app=QtWidgets.QApplication([]); app.setQuitOnLastWindowClosed(False)
window=MainWindow()
for command in window.runtime.list_commands(): window.navigate_to_command(command)
for index in (0,4,5): window.switch_page(index); app.processEvents()
window.execute_task('data.mock', {'count': 2, 'format': 'csv', 'seed': 19, 'template': 'retail_customer'})
deadline=time.monotonic()+30
while window.stack.currentIndex()==2 and time.monotonic()<deadline:
    app.processEvents(); time.sleep(0.01)
assert window.stack.currentIndex()==3
assert all(t['status']=='SUCCEEDED' for t in window.runtime.list_tasks())
QtCore.QThreadPool.globalInstance().waitForDone(30000)
window.close()
print(json.dumps({'status':'passed'}))
''')
            summary["gui"] = gui_result["status"]
        final_hashes = {
            str(path.relative_to(bundled)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in bundled.rglob("*") if path.is_file() and "__pycache__" not in path.parts
        }
        if initial["hashes"] != final_hashes:
            raise AssertionError("Installed execution mutated official plugin resources")
        if (cwd / "workspace").exists() or (module.parent / "workspace").exists():
            raise AssertionError("Installed execution wrote task data into cwd/site-packages")
        return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--python", required=True, type=Path)
    parser.add_argument("--gui", action="store_true")
    parser.add_argument("--evidence", action="store_true")
    options = parser.parse_args()
    print(json.dumps(verify_installed(options.python, gui=options.gui, evidence=options.evidence), indent=2))


if __name__ == "__main__":
    main()
