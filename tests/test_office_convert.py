"""Office acceptance through public Runtime -> real Host -> converter subprocess.

RealLibreOfficeTests generates synthetic *content*, exports genuine OLE XLS/DOC/PPT
with the installed LibreOffice, then verifies the plugin's actual OOXML output.
SyntheticConverterTests instead uses a test-only Python executable. These tests
exercise failure/container protocols, NOT Office conversion or fidelity.

Run with the project's existing environment (no dependency installation):
    .venv/bin/python -m unittest discover -s tests -p test_office_convert.py -v
Missing optional document libraries/soffice skip only the real-engine class.
FODT/FODP use explicit import filters: this alpha misdetects flat ODF as Calc
without them. FODT, not DOCX, produces the table-bearing old DOC: independent
baseline showed DOCX -> MS Word 97 lost the table in fixture production, whereas
FODT -> MS Word 97 -> DOCX retained it. Content assertions are unchanged.
A detected real engine that fails is a FAILURE, never silently downgraded to a
synthetic pass. Failed real fixtures/workspaces are retained in the temp dir.
Windows-native execution is not covered by the POSIX executable fixtures.

Export filter authority (checked 2026-10-07):
https://help.libreoffice.org/latest/en-US/text/shared/guide/convertfilters.html
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
from urllib.parse import unquote, urlparse
from pathlib import Path
from unittest.mock import patch

from testbox.core.runtime import Runtime
from testbox.core.schema_validator import SchemaValidationError

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    from docx import Document
except ImportError:
    Document = None

ROOT = Path(__file__).resolve().parents[1]
PLUGIN = ROOT / "plugins" / "office-convert"
OLE_MAGIC = bytes.fromhex("d0cf11e0a1b11ae1")
ENGINE_ENV = "TESTBOX_OFFICE_CONVERT_SOFFICE_PATH"
REAL_ENGINE = shutil.which("soffice")

# No shell, mocks, private plugin imports, or installed fake application.
# The executable is confined to a disposable test root and selected through
# the same trusted config field as a real converter. Its marker inputs are NOT
# genuine legacy Office; only real tests below make that claim.
FAKE_CONVERTER = r'''
import json, os, subprocess, sys, time, zipfile
from pathlib import Path
from urllib.parse import unquote, urlparse
args = sys.argv[1:]
log = Path(__file__).with_name("converter-calls.jsonl")
record = {"argv": args, "pid": os.getpid()}
profiles = [a for a in args if a.startswith("-env:UserInstallation=")]
if profiles:
    profile = Path(unquote(urlparse(profiles[0].split("=", 1)[1]).path))
    registry = profile / "user" / "registrymodifications.xcu"
    record["registry"] = registry.read_text(encoding="utf-8") if registry.exists() else None
with log.open("a", encoding="utf-8") as stream:
    stream.write(json.dumps(record) + "\n")
if "--version" in args:
    control = Path(__file__).with_name("probe-mode.txt")
    mode = control.read_text().strip() if control.exists() else "valid"
    if mode == "timeout":
        time.sleep(30)
    if mode == "no_version":
        print("SYNTHETIC arbitrary executable with no engine version")
        raise SystemExit(0)
    if mode == "flood":
        os.write(1, b"LibreOffice SYNTHETIC " + b"x" * (1024 * 1024))
        raise SystemExit(0)
    print("LibreOffice SYNTHETIC protocol fixture (not a real engine)")
    raise SystemExit(0)
if "--convert-to" not in args or "--outdir" not in args:
    raise SystemExit(19)
source = Path(args[-1])
mode = source.read_bytes().split(b"mode=")[-1].decode("ascii").strip()
if mode == "timeout_tree":
    child = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    with log.open("a", encoding="utf-8") as stream:
        stream.write(json.dumps({"child_pid": child.pid, "parent_pid": os.getpid()}) + "\n")
    time.sleep(30)
if mode == "timeout":
    time.sleep(30)
if mode in ("stdout_flood", "stderr_flood"):
    fd = 1 if mode == "stdout_flood" else 2
    for _ in range(128):
        os.write(fd, b"SYNTHETIC_DIAGNOSTIC_SECRET " + b"x" * 8192)
    time.sleep(30)
if mode == "exit_nonzero":
    print("SYNTHETIC converter rejection", file=sys.stderr)
    raise SystemExit(23)
if mode == "no_output":
    raise SystemExit(0)
ext = args[args.index("--convert-to") + 1].split(":")[0]
outdir = Path(args[args.index("--outdir") + 1]); outdir.mkdir(parents=True, exist_ok=True)
target = outdir / (source.stem + "." + ext)
if mode == "empty_output":
    target.touch(); raise SystemExit(0)
if mode == "not_zip":
    target.write_bytes(b"SYNTHETIC not OOXML"); raise SystemExit(0)
ctns = "http://schemas.openxmlformats.org/package/2006/content-types"
rns = "http://schemas.openxmlformats.org/package/2006/relationships"
types = {
 "xlsx": ("xl/workbook.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"),
 "docx": ("word/document.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"),
 "pptx": ("ppt/presentation.xml", "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml")}
main, content_type = types[ext]
if mode == "wrong_content_type":
    content_type = "application/octet-stream"
if mode == "macro_content_type":
    content_type = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
parts = {
 "[Content_Types].xml": '<Types xmlns="'+ctns+'"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/'+main+'" ContentType="'+content_type+'"/></Types>',
 "_rels/.rels": '<Relationships xmlns="'+rns+'"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="'+main+'"/></Relationships>'}
if ext == "xlsx":
    parts[main] = '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Synthetic" sheetId="1" r:id="rId1"/></sheets></workbook>'
    parts["xl/_rels/workbook.xml.rels"] = '<Relationships xmlns="'+rns+'"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'
    parts["xl/worksheets/sheet1.xml"] = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>'
elif ext == "docx":
    parts[main] = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>SYNTHETIC</w:t></w:r></w:p></w:body></w:document>'
else:
    parts[main] = '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>'
    parts["ppt/_rels/presentation.xml.rels"] = '<Relationships xmlns="'+rns+'"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/></Relationships>'
    parts["ppt/slides/slide1.xml"] = '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree/></p:cSld></p:sld>'
if mode == "missing_part":
    del parts[main]
if mode == "missing_content_types":
    del parts["[Content_Types].xml"]
if mode == "invalid_xml":
    parts[main] = "<broken>"
if mode == "invalid_content_types_xml":
    parts["[Content_Types].xml"] = "<broken>"
if mode == "traversal":
    parts["../escaped.txt"] = "SYNTHETIC"
if mode == "vba_part":
    parts["xl/vbaProject.bin"] = "SYNTHETIC macro placeholder"
if mode == "activex_part":
    parts["xl/activeX/activeX1.bin"] = "SYNTHETIC ActiveX placeholder"
if mode == "wrong_namespace":
    parts[main] = parts[main].replace("http://schemas.openxmlformats.org/", "urn:synthetic:wrong/")
if mode == "xml_entity":
    parts[main] = '<!DOCTYPE workbook [<!ENTITY synthetic "entity">]>' + parts[main]
with zipfile.ZipFile(target, "w", zipfile.ZIP_STORED) as archive:
    for name, value in parts.items():
        archive.writestr(name, value)
    if mode == "duplicate_part":
        archive.writestr(main, parts[main])
if mode == "bad_crc":
    data = target.read_bytes(); original = parts[main].encode()
    target.write_bytes(data.replace(original, b"!" + original[1:], 1))
'''


ODF_NAMESPACES = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"')


def flat_text_document(paragraphs, rows):
    table_rows = "".join('<table:table-row>' + "".join(
        '<table:table-cell office:value-type="string"><text:p>' + escape(cell) +
        '</text:p></table:table-cell>' for cell in row) + '</table:table-row>' for row in rows)
    return ('<?xml version="1.0" encoding="UTF-8"?><office:document ' + ODF_NAMESPACES +
        ' office:version="1.3" office:mimetype="application/vnd.oasis.opendocument.text">'
        '<office:automatic-styles><style:style style:name="Table1" style:family="table">'
        '<style:table-properties style:width="16cm" table:align="left"/></style:style>'
        '<style:style style:name="Column" style:family="table-column">'
        '<style:table-column-properties style:column-width="8cm"/></style:style></office:automatic-styles>'
        '<office:body><office:text>' + "".join('<text:p>' + escape(text) + '</text:p>' for text in paragraphs) +
        '<table:table table:name="Table1" table:style-name="Table1">'
        '<table:table-column table:style-name="Column" table:number-columns-repeated="2"/>' + table_rows +
        '</table:table></office:text></office:body></office:document>')


def flat_presentation(slides):
    pages = "".join('<draw:page draw:name="Slide' + str(index) +
        '" draw:style-name="DrawingPage" draw:master-page-name="Default">'
        '<draw:frame draw:name="Text' + str(index) +
        '" svg:x="2cm" svg:y="2cm" svg:width="22cm" svg:height="5cm">'
        '<draw:text-box><text:p>' + escape(text) + '</text:p></draw:text-box></draw:frame></draw:page>'
        for index, text in enumerate(slides, 1))
    return ('<?xml version="1.0" encoding="UTF-8"?><office:document ' + ODF_NAMESPACES +
        ' office:version="1.3" office:mimetype="application/vnd.oasis.opendocument.presentation">'
        '<office:automatic-styles><style:page-layout style:name="Page">'
        '<style:page-layout-properties fo:page-width="28cm" fo:page-height="21cm" '
        'style:print-orientation="landscape"/></style:page-layout>'
        '<style:style style:name="DrawingPage" style:family="drawing-page">'
        '<style:drawing-page-properties presentation:background-visible="true" '
        'presentation:background-objects-visible="true"/></style:style></office:automatic-styles>'
        '<office:master-styles><style:master-page style:name="Default" style:page-layout-name="Page" '
        'draw:style-name="DrawingPage"/></office:master-styles><office:body><office:presentation>' + pages +
        '</office:presentation></office:body></office:document>')


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_runtime(root: Path) -> Runtime:
    shutil.copytree(PLUGIN, root / "plugins" / "office-convert",
                    ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    return Runtime(root)


class OfficeTestAssertions:
    """Assertions on public results and persisted artifacts, not plugin internals."""

    def assert_task_persisted(self, task, result):
        self.assertEqual(self.runtime.get_task_result(task), result.to_dict())
        self.assertEqual(self.runtime.get_task(task)["status"],
                         "SUCCEEDED" if result.status == "success" else "FAILED")
        text = self.runtime.get_task_report(task)
        self.assertIn(result.status, text)
        self.assertIn(task, text)
        output = self.root / "workspace" / task / "output"
        for relative in result.files:
            path = output / relative
            self.assertTrue(path.is_file(), result.to_dict())
            path.resolve().relative_to(output.resolve())
        return output

    def assert_staged(self, task, sources, parameter):
        workspace = self.root / "workspace" / task
        manifest = json.loads((workspace / "manifest.json").read_text(encoding="utf-8"))
        records = manifest["inputs"]
        self.assertEqual(len(records), len(sources))
        self.assertEqual(len({item["staged_path"] for item in records}), len(sources))
        for record, source in zip(records, sources):
            self.assertEqual(record["parameter"], parameter)
            self.assertEqual(record["source_path"], str(source.resolve()))
            self.assertEqual(record["size"], source.stat().st_size)
            self.assertEqual(record["sha256"], digest(source))
            staged = workspace / record["staged_path"]
            staged.resolve().relative_to((workspace / "input").resolve())
            self.assertEqual(staged.read_bytes(), source.read_bytes())
        return records

    def conversion_report(self, task, result):
        output = self.assert_task_persisted(task, result)
        reports = [output / path for path in result.files if Path(path).suffix == ".json"]
        self.assertEqual(len(reports), 1, result.to_dict())
        report = json.loads(reports[0].read_text(encoding="utf-8"))
        self.assertIsInstance(report, dict)
        # Public report/Result contract; do not import plugin private helpers.
        self.assertIsInstance(report["items"], list)
        self.assertIsInstance(report["summary"], dict)
        self.assertIn("summary", result.data, result.to_dict())
        self.assertEqual(report["summary"], result.data["summary"])
        return report

    def assert_counts(self, report, *, success, failed, skipped=0):
        summary = report["summary"]
        self.assertEqual(summary["succeeded_count"], success)
        self.assertEqual(summary["failed_count"], failed)
        self.assertEqual(summary["skipped_count"], skipped)
        self.assertEqual(summary["complete"], failed == 0 and skipped == 0)
        self.assertEqual(summary["total_count"], success + failed + skipped)
        self.assertEqual(summary["status"], "succeeded" if not failed and not skipped
                         else "partial" if success else "failed")
        self.assertEqual(len(report["items"]), success + failed + skipped)
        self.assertEqual([item["status"] for item in report["items"]].count("succeeded"), success)
        self.assertEqual([item["status"] for item in report["items"]].count("failed"), failed)
        self.assertEqual([item["status"] for item in report["items"]].count("skipped"), skipped)


@unittest.skipUnless(os.name == "posix", "synthetic executable fixtures are POSIX; Windows native not executed")
class SyntheticConverterTests(OfficeTestAssertions, unittest.TestCase):
    """SYNTHETIC converter protocol only: never Office fidelity acceptance."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="testbox-office-synthetic-"))
        self.addCleanup(self.cleanup_synthetic)
        self.engine = self.root / "synthetic converter with spaces"
        self.engine.write_text("#!" + sys.executable + "\n" + FAKE_CONVERTER, encoding="utf-8")
        self.engine.chmod(0o700)
        self.environment = patch.dict(os.environ, {ENGINE_ENV: str(self.engine)})
        self.environment.start()
        self.addCleanup(self.environment.stop)
        self.runtime = make_runtime(self.root)
        self.addCleanup(self.runtime.close)

    def cleanup_synthetic(self):
        result = self._outcome.result
        if any(test is self or getattr(test, "test_case", None) is self
               for test, _ in result.failures + result.errors):
            print("\nFAILED synthetic protocol artifacts retained at:", self.root, flush=True)
        else:
            shutil.rmtree(self.root)

    def source(self, name="source.xls", mode="valid", payload=None):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload if payload is not None else OLE_MAGIC + b"\nSYNTHETIC protocol input\nmode=" + mode.encode("ascii"))
        return path

    def convert(self, sources, **params):
        inputs = {"input": str(sources)} if isinstance(sources, Path) else {"inputs": list(map(str, sources))}
        before = copy.deepcopy(inputs | params)
        task, result = self.runtime.run("office.convert", inputs | params)
        self.assertEqual(inputs | params, before)
        return task, result

    def calls(self):
        path = self.engine.with_name("converter-calls.jsonl")
        return [json.loads(line) for line in path.read_text().splitlines()] if path.exists() else []

    def test_schema_public_contract_and_defaults(self):
        self.assertEqual(self.runtime.list_unavailable_plugins(), [])
        schema = self.runtime.get_command_schema("office.convert")
        properties = schema["properties"]
        self.assertEqual(properties["input"]["format"], "file-path")
        self.assertEqual(properties["inputs"]["items"]["format"], "file-path")
        self.assertEqual((properties["timeout_seconds"]["minimum"], properties["timeout_seconds"]["maximum"]), (5, 120))
        self.assertEqual((properties["batch_timeout_seconds"]["minimum"], properties["batch_timeout_seconds"]["maximum"]), (10, 240))
        source = self.source()
        params = {"input": str(source)}
        validated = self.runtime.validate_params("office.convert", params)
        self.assertEqual(validated["timeout_seconds"], 60)
        self.assertEqual(validated["batch_timeout_seconds"], 240)
        self.assertIs(validated["continue_on_error"], True)
        self.assertEqual(params, {"input": str(source)})
        self.assertEqual(self.runtime.validate_params("office.inspect", {}), {})
        self.assertEqual(self.runtime.get_command_schema("office.inspect")["properties"], {})

    def test_inspect_synthetic_available_and_mapping(self):
        task, result = self.runtime.run("office.inspect", {})
        self.assertEqual(result.status, "success", result.to_dict())
        self.assertIs(result.data["available"], True)
        self.assertIn("SYNTHETIC", result.data["engine_version"])
        self.assertEqual(result.data["support"], {"xls": "xlsx", "doc": "docx", "ppt": "pptx"})
        self.assert_task_persisted(task, result)

    def test_missing_explicit_engine_inspect_and_convert(self):
        with patch.dict(os.environ, {ENGINE_ENV: str(self.root / "missing-soffice")}):
            task, inspected = self.runtime.run("office.inspect", {})
            self.assertIs(inspected.data["available"], False, inspected.to_dict())
            self.assert_task_persisted(task, inspected)
            task, result = self.convert(self.source())
        self.assertEqual(result.status, "failed", result.to_dict())
        self.assertTrue(result.data.get("error_code"), result.to_dict())
        self.assert_task_persisted(task, result)
        self.assertFalse(any("--convert-to" in call.get("argv", []) for call in self.calls()))

    def test_project_config_and_env_precedence_through_real_host(self):
        # Only a disposable project config is written; no user config is touched.
        (self.root / "config.yaml").write_text('soffice_path: "' + str(self.root / "missing-engine") + '"\n')
        _, available = self.runtime.run("office.inspect", {})
        self.assertIs(available.data["available"], True)
        with patch.dict(os.environ, {ENGINE_ENV: str(self.root / "missing-env-engine")}):
            _, unavailable = self.runtime.run("office.inspect", {})
        self.assertIs(unavailable.data["available"], False)
        (self.root / "config.yaml").write_text('soffice_path: "' + str(self.engine) + '"\n')
        with patch.dict(os.environ):
            os.environ.pop(ENGINE_ENV, None)
            _, available = self.runtime.run("office.inspect", {})
        self.assertIs(available.data["available"], True)

    def test_params_cannot_supply_shell_or_engine_arguments(self):
        source = self.source()
        for key, value in (("soffice_path", str(self.engine)), ("engine_path", str(self.engine)),
                           ("command", "echo unsafe"), ("extra_args", ["--invisible"])):
            with self.subTest(key=key), self.assertRaises(SchemaValidationError):
                self.runtime.run("office.convert", {"input": str(source), key: value})
            with self.subTest(inspect=key), self.assertRaises(SchemaValidationError):
                self.runtime.run("office.inspect", {key: value})
        self.assertEqual(self.calls(), [])

    def test_invalid_types_bounds_unknown_values(self):
        source = self.source()
        cases = [{"input": []}, {"inputs": "bad"}, {"inputs": [42]}, {"inputs": [""]},
                 {"input": str(source), "timeout_seconds": True},
                 {"input": str(source), "timeout_seconds": 4},
                 {"input": str(source), "timeout_seconds": 121},
                 {"input": str(source), "timeout_seconds": 5.0},
                 {"input": str(source), "batch_timeout_seconds": 9},
                 {"input": str(source), "batch_timeout_seconds": 241},
                 {"input": str(source), "continue_on_error": "true"}]
        for params in cases:
            with self.subTest(params=params), self.assertRaises(SchemaValidationError):
                self.runtime.run("office.convert", params)
        self.assertEqual(self.calls(), [])

    def test_no_input_empty_array_and_double_selection_are_rejected(self):
        source = self.source()
        for params in ({}, {"inputs": []}, {"input": ""},
                       {"input": str(source), "inputs": [str(source)]}):
            with self.subTest(params=params):
                try:
                    task, result = self.runtime.run("office.convert", params)
                except (SchemaValidationError, ValueError):
                    continue
                self.assertEqual(result.status, "failed", result.to_dict())
                self.assertTrue(result.data.get("error_code"))
                self.assert_task_persisted(task, result)
        self.assertFalse(any("--convert-to" in call.get("argv", []) for call in self.calls()))

    def test_empty_and_fake_legacy_extensions_do_not_start_converter(self):
        for name, payload in (("empty.xls", b""), ("text.doc", b"not old Office"),
                              ("modern.xls", b"PK\x03\x04not OLE"), ("unsupported.txt", OLE_MAGIC)):
            with self.subTest(name=name):
                source = self.source(name, payload=payload)
                original = source.read_bytes()
                task, result = self.convert(source)
                self.assertEqual(result.status, "failed", result.to_dict())
                report = self.conversion_report(task, result)
                self.assert_counts(report, success=0, failed=1)
                self.assertTrue(report["items"][0]["error_code"])
                self.assertEqual(source.read_bytes(), original)
                self.assertFalse(any(path.startswith("converted/") for path in result.files))
        self.assertFalse(any("--convert-to" in call.get("argv", []) for call in self.calls()))

    def test_synthetic_valid_mapping_staging_and_fixed_argv(self):
        sources = [self.source("input ; $literal.xls"), self.source("a.doc"), self.source("b.ppt")]
        task, result = self.convert(sources)
        self.assertEqual(result.status, "success", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=3, failed=0)
        records = self.assert_staged(task, sources, "inputs")
        converted = [name for name in result.files if name.startswith("converted/")]
        self.assertEqual(converted, [f"converted/{index:04d}-{source.stem}.{ext}"
                                    for index, (source, ext) in enumerate(zip(sources, ("xlsx", "docx", "pptx")), 1)])
        calls = [call for call in self.calls() if "--convert-to" in call.get("argv", [])]
        self.assertEqual(len(calls), 3)
        profiles = []
        workspace = self.root / "workspace" / task
        filters = ("xlsx:Calc MS Excel 2007 XML", "docx:Office Open XML Text",
                   "pptx:Impress MS PowerPoint 2007 XML")
        for call, record, filter_name in zip(calls, records, filters):
            argv = call["argv"]
            self.assertIn("--headless", argv)
            self.assertEqual(Path(argv[-1]).resolve(), (workspace / record["staged_path"]).resolve())
            self.assertEqual(argv[argv.index("--convert-to") + 1], filter_name)
            profile = [arg for arg in argv if arg.startswith("-env:UserInstallation=")]
            self.assertEqual(len(profile), 1)
            profiles.extend(profile)
        for profile in profiles:
            location = Path(unquote(urlparse(profile.split("=", 1)[1]).path))
            location.resolve().relative_to(workspace.resolve())
        self.assertFalse((self.root / "$literal.xlsx").exists())

    def test_synthetic_partial_failure_summary_and_report_preserved(self):
        sources = [self.source("good.xls"), self.source("bad.doc", "no_output"), self.source("last.ppt")]
        task, result = self.convert(sources)
        self.assertEqual(result.status, "success", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=2, failed=1)
        self.assertEqual([item["status"] for item in report["items"]], ["succeeded", "failed", "succeeded"])
        self.assertTrue(report["items"][1]["error_code"])
        self.assertTrue(result.warnings)
        self.assertEqual(len([p for p in result.files if p.startswith("converted/")]), 2)

    def test_synthetic_stop_after_error_skips_unprocessed_items(self):
        sources = [self.source("good.xls"), self.source("bad.doc", "exit_nonzero"), self.source("last.ppt")]
        task, result = self.convert(sources, continue_on_error=False)
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=1, failed=1, skipped=1)
        self.assertEqual([item["status"] for item in report["items"]], ["succeeded", "failed", "skipped"])
        self.assertEqual(len([call for call in self.calls() if "--convert-to" in call.get("argv", [])]), 2)
        self.assertEqual(len([p for p in result.files if p.startswith("converted/")]), 1)

    def test_synthetic_result_summary_equals_registered_report_for_success_partial_and_failure(self):
        cases = (("valid", "valid"), ("valid", "no_output"), ("no_output", "exit_nonzero"))
        for index, modes in enumerate(cases):
            with self.subTest(modes=modes):
                sources = [self.source(f"summary-{index}-{item}.xls", mode)
                           for item, mode in enumerate(modes)]
                task, result = self.convert(sources)
                output = self.assert_task_persisted(task, result)
                registered = [output / name for name in result.files if name.endswith(".json")]
                self.assertEqual(len(registered), 1, result.to_dict())
                report = json.loads(registered[0].read_text(encoding="utf-8"))
                self.assertEqual(result.data["summary"], report["summary"])
                summary = result.data["summary"]
                self.assertEqual(summary["succeeded_count"], modes.count("valid"))
                self.assertEqual(summary["failed_count"], len(modes) - modes.count("valid"))
                self.assertEqual(summary["skipped_count"], 0)
                self.assertIs(summary["complete"], modes == ("valid", "valid"))

    def test_synthetic_all_failure_still_registers_json_and_markdown_reports(self):
        task, result = self.convert([self.source("a.xls", "no_output"), self.source("b.doc", "exit_nonzero")])
        self.assertEqual(result.status, "failed", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=0, failed=2)
        self.assertTrue(any(Path(path).suffix == ".md" for path in result.files))
        self.assertFalse(any(path.startswith("converted/") for path in result.files))

    def test_synthetic_exit_zero_without_output_and_nonzero_exit(self):
        for mode in ("no_output", "empty_output", "exit_nonzero"):
            with self.subTest(mode=mode):
                task, result = self.convert(self.source(mode + ".xls", mode))
                self.assertEqual(result.status, "failed", result.to_dict())
                report = self.conversion_report(task, result)
                self.assert_counts(report, success=0, failed=1)
                self.assertTrue(report["items"][0]["error_code"])
                self.assertFalse(any(path.startswith("converted/") for path in result.files))

    def test_synthetic_invalid_ooxml_rejected_not_registered(self):
        for mode in ("not_zip", "missing_part", "missing_content_types", "wrong_content_type",
                     "invalid_xml", "invalid_content_types_xml", "traversal", "duplicate_part", "bad_crc",
                     "macro_content_type", "vba_part", "activex_part", "xml_entity", "wrong_namespace"):
            with self.subTest(mode=mode):
                task, result = self.convert(self.source(mode + ".xls", mode))
                self.assertEqual(result.status, "failed", result.to_dict())
                report = self.conversion_report(task, result)
                self.assert_counts(report, success=0, failed=1)
                self.assertTrue(report["items"][0]["error_code"])
                self.assertFalse(any(path.startswith("converted/") for path in result.files))
        self.assertFalse((self.root / "escaped.txt").exists())

    def test_synthetic_docx_and_pptx_required_parts_and_content_types(self):
        for extension in ("doc", "ppt"):
            for mode in ("missing_part", "missing_content_types", "wrong_content_type", "invalid_xml"):
                with self.subTest(extension=extension, mode=mode):
                    task, result = self.convert(self.source(mode + "." + extension, mode))
                    self.assertEqual(result.status, "failed", result.to_dict())
                    report = self.conversion_report(task, result)
                    self.assert_counts(report, success=0, failed=1)
                    self.assertTrue(report["items"][0]["error_code"])
                    self.assertFalse(any(path.startswith("converted/") for path in result.files))

    def test_synthetic_single_file_timeout_is_bounded_and_reaped(self):
        before = time.monotonic()
        task, result = self.convert(self.source(mode="timeout"), timeout_seconds=5, batch_timeout_seconds=10)
        elapsed = time.monotonic() - before
        self.assertEqual(result.status, "failed", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=0, failed=1)
        self.assertIn("TIMEOUT", report["items"][0]["error_code"])
        self.assertLess(elapsed, 15, "converter deadline did not bound the subprocess")
        calls = [call for call in self.calls() if "--convert-to" in call.get("argv", [])]
        self.assertEqual(len(calls), 1)
        with self.assertRaises(ProcessLookupError):
            os.kill(calls[0]["pid"], 0)

    def test_synthetic_probe_timeout_missing_version_and_flood_are_not_available(self):
        for mode in ("timeout", "no_version", "flood"):
            with self.subTest(mode=mode):
                self.engine.with_name("probe-mode.txt").write_text(mode)
                before = time.monotonic()
                task, result = self.runtime.run("office.inspect", {})
                self.assertLess(time.monotonic() - before, 12)
                self.assertEqual(result.status, "success", result.to_dict())
                self.assertIs(result.data["available"], False, result.to_dict())
                self.assertIsNone(result.data["engine_version"])
                self.assertTrue(result.warnings)
                self.assert_task_persisted(task, result)

    def test_synthetic_stdout_stderr_flood_bounded_not_persisted(self):
        for mode in ("stdout_flood", "stderr_flood"):
            with self.subTest(mode=mode):
                before = time.monotonic()
                task, result = self.convert(self.source(mode + ".xls", mode), timeout_seconds=5)
                self.assertLess(time.monotonic() - before, 12)
                self.assertEqual(result.status, "failed", result.to_dict())
                report = self.conversion_report(task, result)
                self.assert_counts(report, success=0, failed=1)
                self.assertEqual(report["items"][0]["error_code"], "ENGINE_OUTPUT_LIMIT")
                workspace = self.root / "workspace" / task
                for path in workspace.rglob("*"):
                    if path.is_file():
                        self.assertLess(path.stat().st_size, 1024 * 1024)
                        self.assertNotIn(b"SYNTHETIC_DIAGNOSTIC_SECRET", path.read_bytes())

    def test_synthetic_timeout_owned_wrapper_child_and_profile_cleanup(self):
        task, result = self.convert(self.source(mode="timeout_tree"), timeout_seconds=5)
        self.assertEqual(result.status, "failed", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=0, failed=1)
        self.assertEqual(report["items"][0]["error_code"], "TIMEOUT")
        calls = self.calls()
        wrappers = [call for call in calls if "--convert-to" in call.get("argv", [])]
        children = [call for call in calls if "child_pid" in call]
        self.assertEqual(len(wrappers), 1)
        self.assertEqual(len(children), 1)
        with self.assertRaises(ProcessLookupError):
            os.kill(wrappers[0]["pid"], 0)
        # Probe only the PID this fixture created. No process enumeration,
        # shell, external ps dependency, or termination of unrelated processes.
        child_pid = children[0]["child_pid"]
        deadline = time.monotonic() + 3
        while True:
            try:
                os.kill(child_pid, 0)
            except ProcessLookupError:
                break
            if time.monotonic() >= deadline:
                self.fail("owned converter child survived timeout/reap window")
            time.sleep(0.05)
        record = wrappers[0]
        self.assertIsNotNone(record["registry"])
        self.assertIn("MacroSecurityLevel", record["registry"])
        self.assertIn("DisableMacrosExecution", record["registry"])
        profile_arg = next(arg for arg in record["argv"] if arg.startswith("-env:UserInstallation="))
        profile = Path(unquote(urlparse(profile_arg.split("=", 1)[1]).path))
        profile.resolve().relative_to((self.root / "workspace" / task).resolve())
        self.assertFalse(profile.exists(), "reaped conversion left a live profile behind")
        self.assertFalse(list((self.root / "workspace" / task).rglob("registrymodifications.xcu")))

    def test_synthetic_batch_deadline_marks_remaining_item_skipped(self):
        sources = [self.source("slow1.xls", "timeout"), self.source("slow2.doc", "timeout"), self.source("last.ppt")]
        before = time.monotonic()
        task, result = self.convert(sources, timeout_seconds=120, batch_timeout_seconds=10)
        self.assertLess(time.monotonic() - before, 17)
        self.assertEqual(result.status, "failed", result.to_dict())
        report = self.conversion_report(task, result)
        self.assertIs(report["summary"]["complete"], False)
        self.assertGreaterEqual(report["summary"]["failed_count"], 1)
        self.assertGreaterEqual(report["summary"]["skipped_count"], 1)
        self.assertEqual(report["items"][-1]["status"], "skipped")
        self.assertFalse(any(path.startswith("converted/") for path in result.files))


@unittest.skipUnless(REAL_ENGINE and openpyxl and Document,
                     "real LibreOffice acceptance requires existing soffice/openpyxl/python-docx; nothing is installed")
class RealLibreOfficeTests(OfficeTestAssertions, unittest.TestCase):
    """REAL engine / genuine legacy OLE roundtrips; no fake converter involved."""
    fixture_root = None
    preserve = False

    @classmethod
    def setUpClass(cls):
        cls.fixture_root = Path(tempfile.mkdtemp(prefix="testbox-office-REAL-"))
        cls.addClassCleanup(cls.cleanup_real)
        try:
            version = subprocess.run([REAL_ENGINE, "--version"], capture_output=True, text=True,
                                     timeout=20, check=True)
            cls.engine_version = version.stdout.strip()
            (cls.fixture_root / "engine-version.txt").write_text(cls.engine_version, encoding="utf-8")
            print("\nREAL office engine:", cls.engine_version, "fixtures:", cls.fixture_root, flush=True)
            cls.workbooks = {}
            cls.legacy = {}
            for label, identifier in (("left", "00123"), ("right", "00987")):
                directory = cls.fixture_root / "fixtures" / label
                directory.mkdir(parents=True)
                modern = directory / "同名合成 数据.xlsx"
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["id", "amount", "count", "formula", "notes"])
                sheet.append([identifier, 2, 3, "=SUM(B2:C2)", "中文 Synthetic only"])
                sheet.append(["00000", -2.5, 8, "=B3*C3", "comma, newline\nkept"])
                sheet["A2"].number_format = "@"
                second = workbook.create_sheet("第二张表")
                second.append(["code", "number", "cross-sheet"])
                second.append(["0007", 42, "=Main!B2"])
                workbook.create_sheet("Empty")
                cls.workbooks[label] = {ws.title: [[cell.value for cell in row] for row in ws.iter_rows()]
                                        for ws in workbook.worksheets}
                workbook.save(modern)
                workbook.close()
                cls.legacy[label] = cls.export_legacy(modern, "xls:MS Excel 97")
            document_dir = cls.fixture_root / "fixtures" / "word"
            document_dir.mkdir(parents=True)
            modern = document_dir / "合成正文.docx"
            document = Document()
            cls.paragraphs = ["Synthetic Office acceptance 正文", "Second paragraph: 00123 & <literal>."]
            for text in cls.paragraphs:
                document.add_paragraph(text)
            table = document.add_table(rows=3, cols=2)
            cls.table = [["ID", "Content"], ["00123", "表格 Synthetic row"], ["00000", "last row"]]
            for row, values in zip(table.rows, cls.table):
                for cell, text in zip(row.cells, values):
                    cell.text = text
            document.save(modern)
            flat = document_dir / "合成正文.fodt"
            flat.write_text(flat_text_document(cls.paragraphs, cls.table), encoding="utf-8")
            cls.legacy["word"] = cls.export_legacy(flat, "doc:MS Word 97", "OpenDocument Text Flat XML")
            presentation_dir = cls.fixture_root / "fixtures" / "presentation"
            presentation_dir.mkdir(parents=True)
            cls.slides = ["Synthetic slide one 00123 中文", "Second synthetic slide 00000"]
            flat = presentation_dir / "合成幻灯片.fodp"
            flat.write_text(flat_presentation(cls.slides), encoding="utf-8")
            cls.legacy["presentation"] = cls.export_legacy(
                flat, "ppt:MS PowerPoint 97", "OpenDocument Presentation Flat XML")
        except BaseException:
            cls.preserve = True
            raise

    @classmethod
    def cleanup_real(cls):
        if cls.fixture_root is None:
            return
        if cls.preserve:
            print("\nFAILED real Office artifacts retained at:", cls.fixture_root, flush=True)
        else:
            shutil.rmtree(cls.fixture_root)

    @classmethod
    def export_legacy(cls, modern, filter_name, input_filter=None):
        output = modern.parent / "legacy"
        output.mkdir()
        profile = modern.parent / "export-profile"
        argv = [REAL_ENGINE, "-env:UserInstallation=" + profile.as_uri(), "--headless"]
        if input_filter:
            argv.append("--infilter=" + input_filter)
        argv.extend(["--convert-to", filter_name, "--outdir", str(output), str(modern)])
        process = subprocess.run(argv, capture_output=True, text=True, timeout=60)
        (modern.parent / "fixture-export.json").write_text(json.dumps({"argv": argv,
            "returncode": process.returncode, "stdout": process.stdout, "stderr": process.stderr},
            ensure_ascii=False, indent=2), encoding="utf-8")
        target = output / (modern.stem + "." + filter_name.split(":")[0])
        if process.returncode != 0 or not target.is_file():
            raise AssertionError("REAL fixture export failed: " + json.dumps({"argv": argv,
                                 "returncode": process.returncode, "stdout": process.stdout, "stderr": process.stderr}))
        if target.read_bytes()[:8] != OLE_MAGIC:
            raise AssertionError("REAL legacy export lacks OLE magic: " + str(target))
        return target

    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="task-", dir=self.fixture_root))
        self.environment = patch.dict(os.environ, {ENGINE_ENV: REAL_ENGINE})
        self.environment.start()
        self.addCleanup(self.environment.stop)
        self.runtime = make_runtime(self.root)
        self.addCleanup(self.runtime.close)
        self.originals = {name: (digest(path), path.stat().st_mtime_ns) for name, path in self.legacy.items()}

    def tearDown(self):
        # Retain diagnostic workspaces on failures; successful fixture trees are
        # removed by class cleanup. Neither outcome modifies user Office data.
        result = self._outcome.result
        if any(test is self or getattr(test, "test_case", None) is self
               for test, _ in result.failures + result.errors):
            type(self).preserve = True
        for name, path in self.legacy.items():
            self.assertEqual((digest(path), path.stat().st_mtime_ns), self.originals[name], "source rewritten")

    def assert_workbook(self, output, label):
        workbook = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        expected = self.workbooks[label]
        self.assertEqual(workbook.sheetnames, list(expected))
        for sheet, rows in expected.items():
            actual = [[cell.value for cell in row] for row in workbook[sheet].iter_rows()]
            self.assertEqual(actual, rows, f"Sheet/cell/formula changed: {sheet}")
        self.assertEqual(workbook["Main"]["A2"].data_type, "s")
        self.assertEqual(workbook["Main"]["D2"].data_type, "f")
        values = openpyxl.load_workbook(output, data_only=True)
        self.addCleanup(values.close)
        self.assertEqual(values["Main"]["D2"].value, 5)
        self.assertEqual(values["Main"]["D3"].value, -20)
        self.assertEqual(values["第二张表"]["C2"].value, 2)

    def assert_document(self, output):
        document = Document(output)
        self.assertEqual(len(document.tables), 1,
                         "REAL DOC roundtrip lost the table; text extraction is not table fidelity")
        self.assertEqual([p.text for p in document.paragraphs if p.text], self.paragraphs)
        self.assertEqual([[cell.text for cell in row.cells] for row in document.tables[0].rows], self.table)

    def assert_presentation(self, output):
        # Parse actual OOXML parts; do not infer content from the support map.
        with zipfile.ZipFile(output) as archive:
            self.assertIsNone(archive.testzip())
            root = ET.fromstring(archive.read("ppt/presentation.xml"))
            ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            self.assertEqual(len(root.findall("p:sldIdLst/p:sldId", ns)), len(self.slides))
            names = sorted(name for name in archive.namelist()
                           if name.startswith("ppt/slides/slide") and name.endswith(".xml")
                           and "/_rels/" not in name)
            self.assertEqual(len(names), len(self.slides))
            texts = []
            for name in names:
                slide = ET.fromstring(archive.read(name))
                texts.append("".join(node.text or "" for node in slide.iter()
                                     if node.tag == "{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
            self.assertEqual(texts, self.slides)

    def test_real_inspect_available(self):
        task, result = self.runtime.run("office.inspect", {})
        self.assertEqual(result.status, "success", result.to_dict())
        self.assertIs(result.data["available"], True)
        self.assertIn("LibreOffice", result.data["engine_version"])
        self.assertEqual(result.data["support"], {"xls": "xlsx", "doc": "docx", "ppt": "pptx"})
        self.assert_task_persisted(task, result)

    def test_real_single_xls_all_sheets_leading_zeros_formulas_and_export(self):
        source = self.legacy["left"]
        task, result = self.runtime.run("office.convert", {"input": str(source)})
        self.assertEqual(result.status, "success", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=1, failed=0)
        self.assert_staged(task, [source], "input")
        name = "converted/0001-" + source.stem + ".xlsx"
        self.assertIn(name, result.files)
        output = self.root / "workspace" / task / "output" / name
        self.assert_workbook(output, "left")
        destination = self.root / "user-export" / "chosen.xlsx"
        exported = self.runtime.commit_output(task, name, destination)
        self.assertEqual(exported, destination.resolve())
        self.assertEqual(digest(exported), digest(output))
        self.assertNotEqual(source.resolve(), exported.resolve())

    def test_real_single_doc_body_table_and_staging(self):
        source = self.legacy["word"]
        task, result = self.runtime.run("office.convert", {"input": str(source)})
        self.assertEqual(result.status, "success", result.to_dict())
        name = "converted/0001-" + source.stem + ".docx"
        self.assertIn(name, result.files)
        self.assert_document(self.root / "workspace" / task / "output" / name)
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=1, failed=0)
        self.assert_staged(task, [source], "input")

    def test_real_single_ppt_page_count_and_text_from_fodp_fixture(self):
        source = self.legacy["presentation"]
        task, result = self.runtime.run("office.convert", {"input": str(source)})
        self.assertEqual(result.status, "success", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=1, failed=0)
        self.assert_staged(task, [source], "input")
        name = "converted/0001-" + source.stem + ".pptx"
        self.assertIn(name, result.files)
        self.assert_presentation(self.root / "workspace" / task / "output" / name)

    def test_real_mixed_batch_same_basename_different_sources_do_not_overwrite(self):
        sources = [self.legacy["left"], self.legacy["right"], self.legacy["word"], self.legacy["presentation"]]
        self.assertEqual(sources[0].name, sources[1].name)
        self.assertNotEqual(digest(sources[0]), digest(sources[1]))
        params = {"inputs": list(map(str, sources)), "continue_on_error": True}
        before = copy.deepcopy(params)
        task, result = self.runtime.run("office.convert", params)
        self.assertEqual(params, before)
        self.assertEqual(result.status, "success", result.to_dict())
        report = self.conversion_report(task, result)
        self.assert_counts(report, success=4, failed=0)
        self.assert_staged(task, sources, "inputs")
        names = [f"converted/{index:04d}-{source.stem}.{extension}" for index, (source, extension)
                 in enumerate(zip(sources, ("xlsx", "xlsx", "docx", "pptx")), 1)]
        self.assertEqual([path for path in result.files if path.startswith("converted/")], names)
        output = self.root / "workspace" / task / "output"
        self.assert_workbook(output / names[0], "left")
        self.assert_workbook(output / names[1], "right")
        self.assert_document(output / names[2])
        self.assert_presentation(output / names[3])


if __name__ == "__main__":
    unittest.main()
