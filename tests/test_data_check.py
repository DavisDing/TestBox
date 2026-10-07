"""P3 contract tests using temporary real Runtime/Host tasks and local files."""
from __future__ import annotations

import copy
import csv
import importlib.util
import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from testbox.core.runtime import Runtime
from testbox.core.schema_validator import SchemaValidationError
from testbox.sdk import PluginError, normalize_dataset, read_dataset

ROOT = Path(__file__).resolve().parents[1]
HAS_EXCEL = importlib.util.find_spec("openpyxl") is not None


class DataCheckTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        shutil.copytree(ROOT / "plugins" / "data-check", self.root / "plugins" / "data-check",
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
        self.runtime = Runtime(self.root)
        self.addCleanup(self.runtime.close)

    def source(self, text, suffix=".json", *, encoding="utf-8"):
        path = self.root / ("source" + suffix)
        path.write_text(text, encoding=encoding)
        return path

    def check(self, rows, rules, **kwargs):
        path = self.source(json.dumps(rows, ensure_ascii=False, allow_nan=False))
        return self.run_path(path, rules, **kwargs)

    def run_path(self, path, rules, **kwargs):
        before = path.read_bytes()
        params = {"input": str(path), "rules": rules, **kwargs}
        original = copy.deepcopy(params)
        task_id, result = self.runtime.run("data.check", params)
        self.assertEqual(params, original, "Runtime must not mutate caller parameters")
        self.assertEqual(path.read_bytes(), before, "Plugin must not rewrite source files")
        self.assertEqual(result.status, "success", result.to_dict())
        self.assertEqual(self.runtime.get_task(task_id)["status"], "SUCCEEDED")
        output = self.root / "workspace" / task_id / "output"
        self.assertEqual(result.files, [f"{task_id}.json", f"{task_id}.csv", f"{task_id}.md"])
        report = json.loads((output / f"{task_id}.json").read_text(encoding="utf-8"))
        self.assertEqual(report["task_id"], task_id)
        self.assertEqual(report["passed"], result.data["passed"])
        self.assertNotIn("issues", result.data)  # bounded Host summary
        for key in ("errors", "warnings"):
            self.assertEqual(report[key], sum(rule[key] for rule in report["rule_counts"]))
        self.assertEqual(report["retained_issues"] + report["omitted_issues"], report["issue_count"])
        self.assertEqual(self.runtime.get_task_result(task_id)["status"], "success")
        return report, result, output

    def test_discovery_defaults_and_preview_contract(self):
        self.assertEqual(set(self.runtime.list_commands()), {"data.check"})
        schema = self.runtime.get_command_schema("data.check")
        self.assertEqual(schema["x-preview"], {"command": "data.preview", "sources": [
            {"input": "input", "options": "options", "normalize": "normalize", "label": "输入数据"}]})
        params = self.runtime.validate_params("data.check", {"input": "any.csv", "rules": [{"type": "row_count", "min": 0}]})
        self.assertEqual(params["options"], {})
        self.assertEqual(params["normalize"], {})
        self.assertEqual(params["max_issues"], 1000)
        self.assertEqual(params["rules"][0]["severity"], "error")

    def test_cross_format_real_host(self):
        cases = [
            (".csv", "id,name\n1,Alice\n2,Bob\n", {}),
            (".tsv", "id\tname\n1\tAlice\n2\tBob\n", {}),
            (".txt", "id||name@@1||Alice@@2||Bob", {"delimiter": "||", "record_separator": "@@"}),
            (".json", '{"payload":{"items":[{"id":"1","name":"Alice"},{"id":"2","name":"Bob"}]}}', {"json_path": "$.payload.items"}),
            (".jsonl", '{"id":"1","name":"Alice"}\n{"id":"2","name":"Bob"}\n', {}),
        ]
        rules = [{"type": "required", "fields": ["id", "name"]},
                 {"type": "unique", "fields": ["id"]}, {"type": "type", "field": "id", "value_type": "integer"},
                 {"type": "range", "field": "id", "min": 1, "max": 2}, {"type": "row_count", "min": 2, "max": 2}]
        for suffix, text, options in cases:
            with self.subTest(format=suffix):
                report, _, _ = self.run_path(self.source(text, suffix), rules, options=options, normalize={"types": {"id": "integer"}})
                self.assertTrue(report["passed"])
                self.assertEqual(report["rows_checked"], 2)
                self.assertEqual(report["errors"], 0)

    def test_sql_is_text_not_executed(self):
        sql = "DROP TABLE real_database; SELECT do_not_execute();"
        path = self.source(sql, ".sql")
        report, _, _ = self.run_path(path, [{"type": "required", "fields": ["sql"]},
                                          {"type": "enum", "field": "sql", "values": [sql]}, {"type": "row_count", "min": 1, "max": 1}])
        self.assertTrue(report["passed"])
        self.assertEqual(report["format"], "sql")

    @unittest.skipUnless(HAS_EXCEL, "Excel requires existing optional openpyxl; never install in tests")
    def test_xlsx_xlsm_sheet_headers_and_locations(self):
        from openpyxl import Workbook
        for suffix in (".xlsx", ".xlsm"):
            with self.subTest(suffix=suffix):
                path = self.root / ("source" + suffix)
                book = Workbook()
                book.active.title = "Ignore"
                sheet = book.create_sheet("数据")
                sheet.append(["说明", "说明"])
                sheet.append(["id", "name"])
                sheet.append([1, "Alice"])
                sheet.append([2, ""])
                book.save(path)
                book.close()
                report, _, _ = self.run_path(path, [{"type": "required", "fields": ["name"]}],
                                            options={"sheet": "数据", "header_row": 2, "start_row": 3})
                self.assertEqual(report["errors"], 1)
                self.assertEqual(report["issues"][0]["location"]["row"], 4)
                self.assertEqual(report["issues"][0]["location"]["sheet"], "数据")

    @unittest.skipUnless(HAS_EXCEL, "Excel requires existing optional openpyxl; never install in tests")
    def test_excel_formula_text_default_and_explicit_cached_warning(self):
        from openpyxl import Workbook
        path = self.root / "source.xlsx"
        book = Workbook()
        book.active.append(["value"])
        book.active.append(["=1+1"])
        book.save(path)
        book.close()
        report, _, _ = self.run_path(path, [{"type": "enum", "field": "value", "values": ["=1+1"]}])
        self.assertTrue(report["passed"])
        report, result, _ = self.run_path(path, [{"type": "required", "fields": ["value"]}], options={"formula_mode": "cached", "skip_empty_rows": False})
        self.assertFalse(report["passed"])  # fixture has no formula cache
        self.assertTrue(report["input_warnings"])
        self.assertTrue(result.warnings)

    def test_csv_encoding_quote_header_and_start_row(self):
        path = self.source("说明;说明\r\nid;name\r\n0;忽略\r\n1;'张;三'\r\n", ".csv", encoding="gb18030")
        report, _, _ = self.run_path(path, [{"type": "enum", "field": "name", "values": ["张;三"]}],
                                    options={"encoding": "gb18030", "delimiter": ";", "quotechar": "'", "header_row": 2, "start_row": 4})
        self.assertTrue(report["passed"])
        self.assertEqual(report["rows_checked"], 1)

    def test_headerless_explicit_format_and_columns(self):
        path = self.source("1|Alice\n2|Bob", ".data")
        report, _, _ = self.run_path(path, [{"type": "required", "fields": ["id"]}, {"type": "row_count", "min": 2}],
                                    options={"format": "txt", "delimiter": "|", "has_header": False, "columns": ["id", "name"]})
        self.assertTrue(report["passed"])

    def test_normalize_decimal_exactness_and_source_is_immutable(self):
        raw = [{" ID ": " a ", "amount": " 9007199254740993.1234567890123456789 ", "optional": " NULL "}]
        rules = [{"type": "enum", "field": "id", "values": ["a"]},
                 {"type": "type", "field": "amount", "value_type": "number"},
                 {"type": "range", "field": "amount", "min": "9007199254740993.1234567890123456789", "max": "9007199254740993.1234567890123456789"},
                 {"type": "enum", "field": "optional", "values": [None]}]
        normalize = {"column_mapping": {" ID ": "id"}, "trim": True, "casefold": True,
                     "null_values": ["null"], "types": {"amount": "decimal"}}
        report, _, _ = self.check(raw, rules, normalize=normalize)
        self.assertTrue(report["passed"])
        self.assertEqual(report["types"], {"amount": "decimal"})
        dataset = read_dataset(self.root / "source.json")
        original = copy.deepcopy(dataset)
        normalized = normalize_dataset(dataset, normalize)
        self.assertEqual(dataset, original)
        self.assertEqual(normalized["rows"][0]["amount"], "9007199254740993.1234567890123456789")
        self.assertIsNot(normalized["locations"], dataset["locations"])

    def test_all_explicit_normalized_types_through_real_host(self):
        path = self.source("text,count,amount,flag,day,stamp\nalpha,2,2.50,true,2024-02-29,2024-02-29T01:02:03Z\n", ".csv")
        kinds = {"text": "string", "count": "integer", "amount": "decimal", "flag": "boolean", "day": "date", "stamp": "datetime"}
        rules = [{"type": "type", "field": field, "value_type": "number" if kind == "decimal" else kind}
                 for field, kind in kinds.items()]
        report, _, _ = self.run_path(path, rules, normalize={"types": kinds})
        self.assertTrue(report["passed"])
        self.assertEqual(report["types"], kinds)

    def test_missing_and_null_fail_each_non_nullable_value_rule(self):
        rows = [{"v": 1, "other": 2}, {"v": None, "other": 2}, {"other": 2}]
        rules = [{"type": "type", "field": "v", "value_type": "integer"},
                 {"type": "enum", "field": "v", "values": [1]},
                 {"type": "range", "field": "v", "min": 0},
                 {"type": "compare", "left": "v", "right": "other", "operator": "lt"}]
        report, _, _ = self.check(rows, rules)
        self.assertEqual([rule["errors"] for rule in report["rule_counts"]], [2, 2, 2, 2])
        self.assertEqual(report["errors"], 8)

    def test_required_missing_null_empty_but_not_zero_or_false(self):
        report, _, _ = self.check([{"v": None}, {}, {"v": ""}, {"v": 0}, {"v": False}], [{"type": "required", "fields": ["v"]}])
        self.assertEqual(report["errors"], 3)
        self.assertEqual([issue["location"]["row"] for issue in report["issues"]], [1, 2, 3])
        self.assertFalse(report["passed"])

    def test_absent_column_detected_in_empty_data(self):
        report, _, _ = self.check([], [{"type": "required", "fields": ["missing"]}, {"type": "row_count", "min": 0, "max": 0}])
        self.assertFalse(report["passed"])
        self.assertEqual(report["errors"], 1)
        self.assertEqual(report["issues"][0]["code"], "MISSING_COLUMN")
        self.assertIsNone(report["issues"][0]["location"])

    def test_composite_unique_null_keys_and_bool_distinct_from_int(self):
        rows = [{"a": 1, "b": "x"}, {"a": True, "b": "x"}, {"a": 1.0, "b": "x"},
                {"a": 1, "b": "y"}, {"a": None, "b": "x"}, {"b": "x"}, {"a": "", "b": "x"}]
        report, _, _ = self.check(rows, [{"type": "unique", "fields": ["a", "b"]}])
        self.assertEqual(report["errors"], 4)
        self.assertEqual([issue["code"] for issue in report["issues"]], ["DUPLICATE_KEY", "EMPTY_KEY", "EMPTY_KEY", "EMPTY_KEY"])
        self.assertEqual(report["issues"][0]["location"]["row"], 3)

    def test_unique_json_object_order_and_nested_boolean_equality(self):
        report, _, _ = self.check([{"v": {"a": 1, "b": [True]}}, {"v": {"b": [True], "a": 1.0}}, {"v": {"a": 1, "b": [1]}}], [{"type": "unique", "fields": ["v"]}])
        self.assertEqual(report["errors"], 1)

    def test_strict_type_rules(self):
        cases = [
            ("string", ["x", 1, None], 2), ("integer", [1, True, 1.0, "1"], 3),
            ("number", [1, 1.2, True, "1"], 2), ("boolean", [True, False, 1, "true"], 2),
            ("date", ["2024-02-29", "2025-02-29", "2024-01-01T00:00:00", "20240101"], 3),
            ("datetime", ["2024-02-29T01:02:03+08:00", "2024-02-29", "2025-02-29T01:02:03"], 2),
        ]
        for kind, values, errors in cases:
            with self.subTest(kind=kind):
                report, _, _ = self.check([{"v": value} for value in values], [{"type": "type", "field": "v", "value_type": kind}])
                self.assertEqual(report["errors"], errors)

    def test_enum_strict_bool_and_explicit_null(self):
        report, _, _ = self.check([{"v": True}, {"v": 1}, {"v": None}, {}], [{"type": "enum", "field": "v", "values": [1, None]}])
        self.assertEqual(report["errors"], 2)
        self.assertEqual([issue["location"]["row"] for issue in report["issues"]], [1, 4])

    def test_enum_nested_json_values_through_runtime_schema(self):
        rows = [{"v": {"a": [1, True], "b": None}}, {"v": {"b": None, "a": [1.0, True]}},
                {"v": {"a": [1, 1], "b": None}}]
        report, _, _ = self.check(rows, [{"type": "enum", "field": "v", "values": [{"b": None, "a": [1, True]}]}])
        self.assertEqual(report["errors"], 1)
        self.assertEqual(report["issues"][0]["location"]["row"], 3)

    def test_range_decimal_inclusive_finite_and_boolean(self):
        values = ["1.00000000000000000001", "1.00000000000000000002", "1.00000000000000000003", "NaN", "Infinity", "-Infinity", True, None]
        report, _, _ = self.check([{"v": value} for value in values], [{"type": "range", "field": "v", "min": values[0], "max": values[1]}])
        self.assertEqual(report["errors"], 6)

    def test_length_only_strings_unicode_inclusive(self):
        report, _, _ = self.check([{"v": "你🙂"}, {"v": "a"}, {"v": "abc"}, {"v": [1, 2]}, {"v": 22}, {"v": None}], [{"type": "length", "field": "v", "min": 2, "max": 2}])
        self.assertEqual(report["errors"], 5)

    def test_compare_all_operators_and_strict_types(self):
        for operator, pairs in {"eq": [(1, 1)], "ne": [(True, 1)], "lt": [(1, 2)], "le": [(2, 2)], "gt": [(2, 1)], "ge": [(2, 2)]}.items():
            with self.subTest(operator=operator):
                report, _, _ = self.check([{"a": a, "b": b} for a, b in pairs], [{"type": "compare", "left": "a", "right": "b", "operator": operator}])
                self.assertTrue(report["passed"])
        report, _, _ = self.check([{"a": True, "b": 1}, {"a": None, "b": 1}, {"b": 1}, {"a": 1, "b": "2"}], [{"type": "compare", "left": "a", "right": "b", "operator": "lt"}])
        self.assertEqual(report["errors"], 4)

    def test_decimal_metadata_drives_compare_enum_and_unique(self):
        rows = [{"a": "2.00", "b": "10"}, {"a": "2.0", "b": "10.0"}]
        report, _, _ = self.check(rows, [{"type": "compare", "left": "a", "right": "b", "operator": "lt"},
                                       {"type": "enum", "field": "a", "values": [2]}, {"type": "unique", "fields": ["a"]}],
                                 normalize={"types": {"a": "decimal", "b": "decimal"}})
        self.assertEqual([rule["errors"] for rule in report["rule_counts"]], [0, 0, 1])

    def test_warning_only_and_row_count_boundaries(self):
        report, result, _ = self.check([], [{"type": "row_count", "min": 1, "severity": "warning"}])
        self.assertTrue(report["passed"])
        self.assertEqual(report["errors"], 0)
        self.assertEqual(report["warnings"], 1)
        self.assertTrue(result.warnings)
        report, _, _ = self.check([{"v": 1}], [{"type": "row_count", "min": 1, "max": 1}])
        self.assertTrue(report["passed"])
        report, _, _ = self.check([{"v": 1}, {"v": 2}], [{"type": "row_count", "max": 1}])
        self.assertFalse(report["passed"])

    def test_issue_truncation_never_hides_later_error_or_counts(self):
        rules = [{"type": "required", "fields": ["v"], "severity": "warning"}, {"type": "required", "fields": ["v"]}]
        for maximum in (0, 1, 3, 40):
            with self.subTest(maximum=maximum):
                report, result, output = self.check([{"v": None}] * 20, rules, max_issues=maximum)
                self.assertFalse(report["passed"])
                self.assertEqual((report["errors"], report["warnings"]), (20, 20))
                self.assertEqual(report["retained_issues"], maximum)
                self.assertEqual(report["truncated"], maximum < 40)
                self.assertEqual([rule["checked"] for rule in report["rule_counts"]], [20, 20])
                self.assertIn(f"truncated：{str(maximum < 40).lower()}", (output / result.files[2]).read_text(encoding="utf-8"))
                csv_rows = list(csv.DictReader(io.StringIO((output / result.files[1]).read_text(encoding="utf-8-sig"))))
                self.assertEqual(csv_rows[0]["record_type"], "summary")
                self.assertEqual(csv_rows[0]["errors"], "20")
                self.assertEqual(csv_rows[0]["truncated"], str(maximum < 40))
                self.assertEqual(sum(row["record_type"] == "issue" for row in csv_rows), maximum)
                if maximum == 1:
                    self.assertEqual(report["issues"][0]["severity"], "warning")

    def test_locations_keep_multiline_csv_physical_rows(self):
        path = self.source('id,v\n1,"line1\nline2"\n2,\n', ".csv")
        report, _, _ = self.run_path(path, [{"type": "required", "fields": ["v"]}])
        self.assertEqual(report["issues"][0]["location"]["row"], 4)
        self.assertTrue(report["issues"][0]["location"]["source"].endswith("source.csv"))

    def test_bad_configuration_never_silently_ignored(self):
        bad_rules = [
            [], [{"type": "unknown"}], [{"type": "regex", "field": "v", "pattern": "(a+)+$"}],
            [{"type": "required", "fields": []}], [{"type": "required", "fields": ["v", "v"]}],
            [{"type": "unique", "fields": [""]}], [{"type": "type", "field": "v", "value_type": "bytes"}],
            [{"type": "enum", "field": "v", "values": []}], [{"type": "range", "field": "v"}],
            [{"type": "range", "field": "v", "min": "NaN"}], [{"type": "range", "field": "v", "max": "Infinity"}],
            [{"type": "range", "field": "v", "min": True}], [{"type": "range", "field": "v", "min": 2, "max": 1}],
            [{"type": "length", "field": "v", "min": -1}], [{"type": "length", "field": "v", "max": 2.5}],
            [{"type": "row_count", "min": True}], [{"type": "compare", "left": "v", "right": "x", "operator": "eval"}],
            [{"type": "required", "fields": ["v"], "severity": "ignore"}],
            [{"type": "required", "fields": ["v"], "expression": "__import__('os')"}],
            [{"type": "range", "field": "v", "min": 0, "fields": ["v"]}],
            [{"type": "row_count", "min": 0, "id": "x"}, {"type": "row_count", "min": 0, "id": "x"}],
        ]
        path = self.source("not JSON")  # bad rule validation must run before reader
        for rules in bad_rules:
            with self.subTest(rules=rules):
                try:
                    task_id, result = self.runtime.run("data.check", {"input": str(path), "rules": rules})
                except SchemaValidationError:
                    continue
                self.assertEqual(result.status, "failed")
                self.assertEqual(result.data["error_code"], "INVALID_PARAMS", result.to_dict())
                self.assertEqual(result.files, [])
                self.assertEqual(self.runtime.get_task(task_id)["status"], "FAILED")

    def test_bad_top_level_and_sdk_configuration(self):
        path = self.source('[{"v":1}]')
        base = {"input": str(path), "rules": [{"type": "row_count", "min": 0}]}
        for extra in ({"max_issues": True}, {"max_issues": -1}, {"max_issues": 10001}, {"options": []}, {"normalize": []},
                      {"options": {"unknown": True}}, {"normalize": {"unknown": True}}, {"options": {"delimiter": 3}},
                      {"normalize": {"types": {"missing": "integer"}}}, {"options": {"max_rows": True}}):
            with self.subTest(extra=extra):
                try:
                    _, result = self.runtime.run("data.check", {**base, **extra})
                except SchemaValidationError:
                    continue
                self.assertEqual(result.status, "failed", result.to_dict())
                self.assertEqual(result.files, [])
                self.assertIn(result.data["error_code"], {"INVALID_PARAMS", "INPUT_INVALID"})

    def test_reader_limits_duplicate_json_keys_and_nonfinite_data_fail_task(self):
        for text, options in [('[{"v":1},{"v":2}]', {"max_rows": 1}), ('[{"v":1,"v":2}]', {}),
                              ('[{"v":NaN}]', {}), ('[{"v":Infinity}]', {}), ('[{"v":1e999}]', {})]:
            with self.subTest(text=text):
                path = self.source(text)
                _, result = self.runtime.run("data.check", {"input": str(path), "options": options, "rules": [{"type": "row_count", "min": 0}]})
                self.assertEqual(result.status, "failed")
                self.assertEqual(result.files, [])
                self.assertNotIn("passed", result.data)
        self.assertEqual(result.data["error_code"], "INPUT_INVALID")

    def test_no_original_values_in_logs_reports_and_formula_safe_csv(self):
        secret = "private-value-do-not-log"
        report, result, output = self.check([{"v": secret}], [{"type": "range", "field": "v", "min": 0, "id": "=HYPERLINK(\"evil\")"}])
        self.assertFalse(report["passed"])
        for name in result.files:
            self.assertNotIn(secret, (output / name).read_text(encoding="utf-8-sig"))
        task_id = report["task_id"]
        self.assertNotIn(secret, self.runtime.get_task_log(task_id))
        parsed = list(csv.DictReader(io.StringIO((output / result.files[1]).read_text(encoding="utf-8-sig"))))
        self.assertEqual(next(row for row in parsed if row["record_type"] == "issue")["rule_id"], "'=HYPERLINK(\"evil\")")

    def test_normalization_failure_has_no_original_values_in_log(self):
        secret = "private-failed-conversion"
        path = self.source(json.dumps([{"v": secret}]))
        before = path.read_bytes()
        task_id, result = self.runtime.run("data.check", {"input": str(path), "rules": [{"type": "row_count", "min": 0}], "normalize": {"types": {"v": "integer"}}})
        self.assertEqual(result.status, "failed")
        self.assertNotIn(secret, self.runtime.get_task_log(task_id))
        self.assertNotIn(secret, json.dumps(result.to_dict()))
        self.assertEqual(path.read_bytes(), before)

    def test_incomplete_sdk_contract_refused_and_internal_rule_validation(self):
        # No reader mocks: this unit check additionally protects against a future
        # SDK contract changing complete=true to an implicit sampled result.
        spec = importlib.util.spec_from_file_location("data_check_unit", ROOT / "plugins/data-check/src/main.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        rules = module.validate_params({"input": "file.csv", "rules": [{"type": "row_count", "min": 0}]})
        with self.assertRaises(PluginError) as caught:
            module.check_dataset({"rows": [], "locations": [], "columns": [], "complete": False}, rules, 0)
        self.assertEqual(caught.exception.code, "INCOMPLETE_DATASET")
        for value in (True, float("nan"), float("inf"), "NaN", "-Infinity"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                module.numeric(value)
        for text in ("=x", "+x", "-x", "@x", " \t=x", "\tx", "\rx", "\nx"):
            self.assertTrue(module.csv_safe(text).startswith("'"))


if __name__ == "__main__":
    unittest.main()
