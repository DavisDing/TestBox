"""Real Runtime/Host contracts for fixed records and independent batch units."""
import copy
import json
from pathlib import Path
import shutil
import tempfile
import unittest

from testbox.core.runtime import Runtime
from testbox.sdk import PluginError, read_dataset
from testbox.dataset_batch import expand

try:
    import openpyxl
except ImportError:
    openpyxl = None
ROOT = Path(__file__).resolve().parents[1]


class FixedWidthTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def read(self, content, **options):
        p = self.root / 'fixed.txt'; p.write_bytes(content)
        return read_dataset(p, {'format': 'fixed', 'widths': [3, 2], 'has_header': False,
                                'columns': ['id', 'name'], **options})

    def test_characters_preserve_padding_zero_and_location(self):
        data = self.read('001张三\r\n002李 \r\n'.encode())
        self.assertEqual(data['rows'], [{'id':'001','name':'张三'}, {'id':'002','name':'李 '}])
        self.assertEqual([l['row'] for l in data['locations']], [1,2])
        self.assertTrue(data['complete'])

    def test_bytes_gbk_utf8_bom_and_split_character_refusal(self):
        data = self.read('001张三\n'.encode('gbk'), encoding='gbk', widths=[3,4], width_unit='bytes')
        self.assertEqual(data['rows'][0]['name'], '张三')
        data = self.read(b'\xef\xbb\xbf'+ '001张三\n'.encode(), widths=[3,6], width_unit='bytes')
        self.assertEqual(data['rows'][0]['name'], '张三')
        with self.assertRaises(PluginError):
            self.read('001张三'.encode(), widths=[4,5], width_unit='bytes')

    def test_custom_separator_header_and_empty_records(self):
        data = self.read(b'id v <EOR>001AB<EOR>002CD<EOR>', has_header=True,
                         record_separator='<EOR>')
        self.assertEqual(data['columns'], ['id ', 'v '])
        self.assertEqual(data['rows'][0], {'id ': '001', 'v ': 'AB'})
        # Header is also a strict-width record.

    def test_short_long_invalid_widths_and_resource_limits_fail(self):
        for value in (b'001A', b'001ABC'):
            with self.subTest(value=value), self.assertRaises(PluginError): self.read(value)
        for widths in ([], [True,2], [0,2], [-1,2], [1.1,2], '3,2'):
            with self.subTest(widths=widths), self.assertRaises(PluginError): self.read(b'001AB', widths=widths)
        with self.assertRaises(PluginError): self.read(b'001AB\n002CD', max_rows=1)
        with self.assertRaises(PluginError): self.read(b'001AB', max_bytes=2)
        with self.assertRaises(PluginError): self.read(b'001AB', width_unit='unknown')


class BatchHostTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        shutil.copytree(ROOT/'plugins',self.root/'plugins',ignore=shutil.ignore_patterns('__pycache__'))
        self.runtime=Runtime(self.root);self.addCleanup(self.runtime.close)

    def file(self, directory, name, text):
        root=self.root/directory;root.mkdir(exist_ok=True)
        path=root/name;path.write_text(text, encoding='utf-8');return str(path)

    def run_batch(self,command,params):
        before=copy.deepcopy(params)
        sources={p:Path(p).read_bytes() for key, values in params.items() if key in ('inputs','left_inputs','right_inputs') for p in values}
        task,result=self.runtime.run(command,params)
        self.assertEqual(params,before)
        for path,content in sources.items():self.assertEqual(Path(path).read_bytes(),content)
        report=json.loads((self.root/'workspace'/task/'output'/f'{task}-batch.json').read_text())
        for name in result.files:self.assertTrue((self.root/'workspace'/task/'output'/name).is_file(),name)
        return task,result,report

    def test_preview_partial_failure_retains_reports_and_task_history(self):
        good=self.file('left','a.csv','id,v\n001,A\n');bad=self.file('left','b.csv','id,id\n1,2\n')
        task,result,report=self.run_batch('data.preview',{'inputs':[bad,good]})
        self.assertEqual(result.status,'failed');self.assertEqual(report['succeeded'],1)
        display_data=json.loads((self.root/'workspace'/task/'output'/report['items'][1]['preview_file']).read_text())
        self.assertEqual(display_data['rows'][0]['id'],'001')
        self.assertEqual(self.runtime.get_task(task)['status'],'FAILED')
        display=report['items'][1]['preview_file']
        self.assertTrue(self.runtime.get_task_artifact_path(task,display).is_file())
        with self.assertRaises(ValueError):self.runtime.get_task_output_path(task,display)
        with self.assertRaises(ValueError):self.runtime.get_task_artifact_path(task,'../result.json')

    def test_name_pair_cross_format_unmatched_and_no_false_equal(self):
        a=self.file('left','a.csv','id\n001\n');b=self.file('left','b.csv','id\n002\n')
        ar=self.file('right','a.json','[{"id":"001"}]')
        _,result,report=self.run_batch('data.compare',{'left_inputs':[b,a],'right_inputs':[ar]})
        self.assertEqual(result.status,'failed');self.assertFalse(result.data['equal'])
        self.assertEqual(report['items'][0]['status'],'unmatched')
        self.assertTrue(report['items'][1]['data']['equal'])

    def test_equal_and_different_are_business_outcomes(self):
        a=self.file('left','a.csv','id\n001\n');b=self.file('left','b.csv','id\n002\n')
        ar=self.file('right','a.csv','id\n001\n');br=self.file('right','b.csv','id\n003\n')
        _,result,report=self.run_batch('data.compare',{'left_inputs':[b,a],'right_inputs':[ar,br]})
        self.assertEqual(result.status,'success');self.assertFalse(report['equal'])
        self.assertEqual(report['succeeded'],2)

    def test_ambiguous_names_not_arbitrarily_paired(self):
        a=self.file('left','a.csv','id\n1\n');b=self.file('left','a.json','[{"id":"1"}]')
        ar=self.file('right','a.csv','id\n1\n')
        _,result,report=self.run_batch('data.compare',{'left_inputs':[a,b],'right_inputs':[ar]})
        self.assertEqual(result.status,'failed')
        self.assertTrue(all(i.get('code')=='AMBIGUOUS_PAIR' for i in report['items']))

    def test_position_pairing_and_quality_rule_counts_independent(self):
        a=self.file('left','one.csv','id\n001\n');b=self.file('left','two.csv','id\n001\n')
        ar=self.file('right','different.csv','id\n001\n')
        _,result,report=self.run_batch('data.compare',{'left_inputs':[a],'right_inputs':[ar],'batch':{'pairing':'position'}})
        self.assertTrue(report['equal'])
        _,result,report=self.run_batch('data.check',{'inputs':[a,b],'rules':[{'type':'unique','fields':['id']}]})
        self.assertTrue(report['passed']);self.assertEqual(report['succeeded'],2)

    def test_fixed_through_all_three_plugins(self):
        a=self.file('left','a.txt','001AB\n002CD\n');b=self.file('right','b.txt','001AB\n002CD\n')
        options={'format':'fixed','widths':[3,2],'has_header':False,'columns':['id','v']}
        _,r=self.runtime.run('data.preview',{'input':a,'options':options});self.assertEqual(r.status,'success',r.to_dict())
        _,r=self.runtime.run('data.compare',{'left':a,'right':b,'left_options':options,'right_options':options});self.assertTrue(r.data['equal'])
        _,r=self.runtime.run('data.check',{'input':a,'options':options,'rules':[{'type':'unique','fields':['id']}]});self.assertTrue(r.data['passed'])

    def test_conflicts_and_batch_options_refused(self):
        a=self.file('left','a.csv','id\n001\n')
        for extra in ({'input':a}, {'batch':{'unknown':True}}, {'options':{'sheets':'all'}}, {'batch':{'pairing':'oops'}}):
            _,r=self.runtime.run('data.preview',{'inputs':[a],**extra})
            self.assertEqual(r.status,'failed',r.to_dict())
        with self.assertRaises(PluginError):expand([a]*101,{})

    def test_sql_batch_does_not_report_unknown_equal(self):
        a=self.file('left','a.sql','SELECT * FROM t;');b=self.file('right','b.sql','SELECT * FROM t;')
        _,r,report=self.run_batch('sql.diff',{'left_inputs':[a],'right_inputs':[b]})
        self.assertEqual(r.status,'success');self.assertEqual(report['verdict'],'inconclusive')

    def test_batch_expansion_limits_and_null_selection(self):
        paths=[str(self.root/f"file-{i}.csv") for i in range(101)]
        with self.assertRaises(PluginError): expand(paths,{})
        with self.assertRaises(PluginError): expand([],{'sheets':None})
        a=self.file('left','a.csv','id\n1\n')
        from testbox.dataset_batch import run_batch
        for params in ({'inputs':[a],'batch':{'pairing':[]}}, {'inputs':[a,a]}):
            with self.assertRaises(PluginError):run_batch(None,'data.preview',params)

    @unittest.skipUnless(openpyxl,'openpyxl not installed')
    def test_too_many_sheets_and_duplicate_selection_fail_closed(self):
        path=self.root/'many.xlsx';book=openpyxl.Workbook()
        for index in range(100):book.create_sheet(str(index))
        book.save(path);book.close()
        units=expand([str(path)],{'sheets':'all'})
        self.assertEqual(units[0]['failure'].code,'INPUT_LIMIT_EXCEEDED')
        with self.assertRaises(PluginError):expand([str(path)],{'sheets':['Sheet','Sheet']})

    @unittest.skipUnless(openpyxl,'openpyxl not installed')
    def test_all_sheets_mapping_missing_and_per_sheet_failure(self):
        paths=[]
        for directory,names in [('left',['Customers','Broken','Orders']),('right',['Clients','Broken','Orders'])]:
            folder=self.root/directory;folder.mkdir()
            path=folder/'data.xlsx';book=openpyxl.Workbook();book.remove(book.active)
            for name in names:
                sheet=book.create_sheet(name);sheet.append(['id','id'] if name=='Broken' else ['id','v']);sheet.append(['001','A'])
            book.save(path);book.close();paths.append(str(path))
        _,r,report=self.run_batch('data.compare',{'left_inputs':[paths[0]],'right_inputs':[paths[1]],
          'left_options':{'sheets':'all'},'right_options':{'sheets':'all'},'batch':{'sheet_mapping':{'Customers':'Clients'}}})
        self.assertEqual(r.status,'failed');self.assertEqual(report['succeeded'],2);self.assertFalse(report['equal'])
        task,r,report=self.run_batch('data.preview',{'inputs':[paths[0]],'options':{'sheets':['Orders','Missing','Customers']}})
        self.assertEqual(report['succeeded'],2);self.assertEqual([i['input']['sheet'] for i in report['items']],['Orders','Missing','Customers'])
        self.assertTrue(self.runtime.get_task_artifact_path(task,report['items'][0]['preview_file']).is_file())

if __name__=='__main__':unittest.main()
