from __future__ import annotations

import copy
import importlib.util
import json
from pathlib import Path
import shutil
import sqlite3
import tempfile
import unittest

from testbox.core.runtime import Runtime
from testbox.sdk import PluginError, read_dataset


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("codecheck_helpers", ROOT / "plugins/code-value-check/src/main.py")
HELPERS = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(HELPERS)


class CodeValueCheckTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        shutil.copytree(ROOT / "plugins/code-value-check", self.root / "plugins/code-value-check")
        self.runtime = Runtime(self.root)
        self.addCleanup(self.runtime.close)
        self.plan_columns = {"source_table": "src", "source_code_field": "src_field", "target_table": "dst",
                             "target_code_field": "dst_field", "domain_value": "domain"}
        self.mapping_columns = {"domain": "domain", "source_code": "src", "target_code": "dst", "code_label": "label"}
        self.plan = self.file("plan.json", [{"src": "a", "src_field": "code", "dst": "b", "dst_field": "code", "domain": "GENDER"}])
        self.params = {"test_plan_input": str(self.plan), "test_plan_columns": self.plan_columns,
                       "mapping_mode": "database", "sql_engine": "sqlite", "engine_version": "3.35",
                       "defaults": {"source_keys": ["id"], "target_keys": ["id"]},
                       "code_table": {"table_parts": ["mapping"], "columns": self.mapping_columns}}

    def file(self, name, value):
        path = self.root / name
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def generate(self, **changes):
        params = copy.deepcopy(self.params)
        params.update(changes)
        task, result = self.runtime.run("sql.codecheck", params)
        self.assertEqual(result.status, "success", result.to_dict())
        script = self.runtime.get_task_output_path(task, "codecheck.sql").read_text(encoding="utf-8")
        plan = json.loads(self.runtime.get_task_output_path(task, "codecheck-plan.json").read_text(encoding="utf-8"))
        self.assertEqual(result.data["execution_status"], "NOT_EXECUTED")
        return script, plan, result

    def database(self, sources=(), targets=(), mappings=()):
        database = sqlite3.connect(":memory:")
        self.addCleanup(database.close)
        database.row_factory = sqlite3.Row
        database.executescript("CREATE TABLE a(id TEXT, code TEXT, region TEXT, enabled TEXT);"
                               "CREATE TABLE b(id TEXT, code TEXT, region TEXT, enabled TEXT);"
                               "CREATE TABLE mapping(domain TEXT, src TEXT, dst TEXT, label TEXT, region TEXT, system TEXT, enabled TEXT, source_table TEXT);")
        database.executemany("INSERT INTO a VALUES (?,?,?,?)", [tuple(row) + (None,) * (4 - len(row)) for row in sources])
        database.executemany("INSERT INTO b VALUES (?,?,?,?)", [tuple(row) + (None,) * (4 - len(row)) for row in targets])
        database.executemany("INSERT INTO mapping VALUES (?,?,?,?,?,?,?,?)", [tuple(row) + (None,) * (8 - len(row)) for row in mappings])
        return database

    def execute(self, database, script):
        results, statement = [], ""
        for line in script.splitlines(keepends=True):
            statement += line
            if sqlite3.complete_statement(statement):
                cursor = database.execute(statement)
                if cursor.description:
                    results.append([dict(row) for row in cursor.fetchall()])
                statement = ""
        self.assertFalse(statement.strip())
        return results

    def test_real_host_six_modes_keep_duplicates_and_domain_isolation(self):
        rows = [{"src": "a", "src_field": "code", "dst": "b", "dst_field": "code", "domain": "GENDER", "case_id": "first"},
                {"src": "a", "src_field": "code", "dst": "b", "dst_field": "code", "domain": "GENDER", "case_id": "second"}]
        self.plan.write_text(json.dumps(rows), encoding="utf-8")
        self.params["test_plan_columns"]["case_id"] = "case_id"
        for mode in ("select_only", "insert_only", "select_and_insert"):
            for layout in ("separate", "union_all"):
                with self.subTest(mode=mode, layout=layout):
                    script, plan, result = self.generate(output_mode=mode, statement_layout=layout, include_summary=True)
                    database = self.database(sources=[("ok", "1"), ("bad", "1"), ("dup-src", "1"), ("dup-src", "2"), ("dup-dst", "1")],
                                             targets=[("ok", "a"), ("bad", "b"), ("missing", "a"), ("dup-src", "a"), ("dup-dst", "a"), ("dup-dst", "a")],
                                             mappings=[("GENDER", "1", "a", "男"), ("GENDER", "2", "b", "女"), ("ID", "1", "idcard", "身份证")])
                    results = self.execute(database, script)
                    if mode == "insert_only":
                        detail = [dict(row) for row in database.execute('SELECT * FROM "' + plan["result_table"][0] + '"')]
                    else:
                        detail = [row for group in results for row in group if "check_status" in row]
                    self.assertEqual(len(detail), 12)
                    for case_id in ("first", "second"):
                        selected = [row for row in detail if row["case_id"] == case_id]
                        self.assertCountEqual([row["check_status"] for row in selected], ["SAME", "DIFFERENT", "SOURCE_MISSING", "DUPLICATE_SOURCE_MATCH", "DUPLICATE_TARGET_KEY", "DUPLICATE_TARGET_KEY"])
                        duplicate = next(row for row in selected if row["check_status"] == "DUPLICATE_SOURCE_MATCH")
                        self.assertIsNone(duplicate["source_key_1"])
                        self.assertIsNone(duplicate["expected_target_code"])
                    summary = [row for group in results for row in group if "total_rows" in row]
                    self.assertEqual(len(summary), 2)
                    self.assertTrue(all(row["total_rows"] == 6 and row["same_rows"] == 1 and row["not_comparable_rows"] == 4 for row in summary))
                    self.assertTrue(all(abs(row["pass_rate"] - 1 / 6) < 1e-9 for row in summary))
                    self.assertTrue(script.startswith("CREATE TABLE") if mode != "select_only" else script.startswith("WITH"))

    def test_null_duplicate_missing_and_empty_keys(self):
        script, _, _ = self.generate()
        scenarios = [
            ([("1", "1")], [("1", "a")], [("ID", "1", "a", None)], "MAPPING_MISSING", 0),
            ([("1", "1")], [("1", "a")], [("GENDER", "1", "a", None)] * 2, "DUPLICATE_MAPPING", 2),
            ([("1", "1")], [("1", "a")], [("GENDER", "1", "a", None), ("GENDER", "1", "b", None)], "DUPLICATE_MAPPING", 2),
            ([("1", None)], [("1", "a")], [], "SOURCE_CODE_NULL", None),
            ([("1", "1")], [("1", None)], [("GENDER", "1", None, None)], "NULL_VALUE", 1),
            ([("1", "1")], [("1", "a")], [("GENDER", "1", None, None)], "DIFFERENT", 1),
            ([(None, "1")], [(None, "a")], [], "INVALID_TARGET_KEY", None),
            ([("", "1")], [("", "a")], [], "INVALID_TARGET_KEY", 0),
            ([("1", "01")], [("1", "a")], [("GENDER", "1", "a", None)], "MAPPING_MISSING", 0),
        ]
        for sources, targets, mappings, status, count in scenarios:
            with self.subTest(status=status, mappings=mappings):
                row = self.execute(self.database(sources, targets, mappings), script)[0][0]
                self.assertEqual(row["check_status"], status)
                self.assertEqual(row["mapping_match_count"], count)
                if count != 1:
                    self.assertIsNone(row["expected_target_code"])
        params = copy.deepcopy(self.params["defaults"])
        params["nulls_equal"] = True
        null_script, _, _ = self.generate(defaults=params)
        row = self.execute(self.database([("1", "1")], [("1", None)], [("GENDER", "1", None, None)]), null_script)[0][0]
        self.assertEqual(row["check_status"], "SAME")

    def test_dynamic_and_static_conditions_preserve_counts_and_evidence(self):
        conditions = [{"mapping_field": "region", "operator": "eq", "value_source": "source_field", "field": "region"},
                      {"mapping_field": "system", "operator": "eq", "value": "CRM"},
                      {"mapping_field": "source_table", "operator": "eq", "value_source": "case_field", "field": "source_table"},
                      {"mapping_field": "enabled", "operator": "in", "value": ["Y", "YES"]}]
        defaults = {**self.params["defaults"], "mapping_conditions": conditions,
                    "source_filters": [{"field": "enabled", "operator": "eq", "value": "Y"}],
                    "target_filters": [{"field": "enabled", "operator": "eq", "value": "Y"}]}
        columns = {**self.mapping_columns, **{name: name for name in ("region", "system", "source_table", "enabled")}}
        script, _, _ = self.generate(defaults=defaults, code_table={"table_parts": ["mapping"], "columns": columns})
        base_mapping = ("GENDER", "1", "a", "男", "north", "CRM", "Y", "a")
        others = [("ID", "1", "x", "身份证", "north", "CRM", "Y", "a"),
                  ("GENDER", "1", "b", "女", "south", "CRM", "Y", "a"),
                  ("GENDER", "1", "x", "停用", "north", "CRM", "N", "a"),
                  ("GENDER", "1", "x", "系统", "north", "OTHER", "Y", "a"),
                  ("GENDER", "1", "x", "表", "north", "CRM", "Y", "other")]
        for repeat, status in ((0, "MAPPING_MISSING"), (1, "SAME"), (2, "DUPLICATE_MAPPING")):
            database = self.database([("1", "1", "north", "Y"), ("1", "2", "south", "N")],
                                     [("1", "a", None, "Y"), ("2", "a", None, "N")], others + [base_mapping] * repeat)
            row = self.execute(database, script)[0][0]
            self.assertEqual(row["check_status"], status)
            self.assertEqual(row["mapping_match_count"], repeat)
            self.assertEqual(row["dimension_value_1"], "north")
            self.assertEqual(row["dimension_source_field_1"], "region")

    def test_reverse_and_zero_data_summary(self):
        script, _, _ = self.generate(include_summary=True, include_reverse_check=True)
        database = self.database(sources=[("missing", "1"), (None, "1"), ("dup", "1"), ("dup", "1")])
        results = self.execute(database, script)
        details = results[0]
        self.assertCountEqual([row["check_status"] for row in details], ["TARGET_MISSING", "INVALID_SOURCE_KEY", "DUPLICATE_SOURCE_KEY", "DUPLICATE_SOURCE_KEY"])
        summary = results[1]
        forward = next(row for row in summary if row["check_direction"] == "target_to_source")
        reverse = next(row for row in summary if row["check_direction"] == "source_to_target")
        self.assertEqual(forward["data_status"], "NO_DATA")
        self.assertEqual(forward["total_rows"], 0)
        self.assertIsNone(forward["pass_rate"])
        self.assertEqual(reverse["total_rows"], 4)
        self.assertIsNone(reverse["pass_rate"])

    def test_errors_scope_and_stored_summary(self):
        for mode in ("select_only", "insert_only", "select_and_insert"):
            changes = {"output_mode": mode, "include_summary": True}
            if mode != "insert_only":
                changes["result_scope"] = "errors_only"
            if mode == "insert_only":
                changes["store_scope"] = "errors_only"
            script, plan, _ = self.generate(**changes)
            database = self.database([("ok", "1"), ("bad", "1")], [("ok", "a"), ("bad", "b")], [("GENDER", "1", "a", "男")])
            results = self.execute(database, script)
            summary = [row for group in results for row in group if "total_rows" in row]
            self.assertEqual(summary[0]["total_rows"], 2)
            details = ([dict(row) for row in database.execute('SELECT * FROM "' + plan["result_table"][0] + '"')] if mode == "insert_only"
                       else [row for group in results for row in group if "check_status" in row])
            self.assertEqual(len(details), 1)
            self.assertEqual(details[0]["check_status"], "DIFFERENT")

    def test_composite_keys_different_names_and_padded_union_shape(self):
        defaults = {"source_keys": ["id", "region"], "target_keys": ["id", "region"]}
        script, _, _ = self.generate(defaults=defaults, statement_layout="union_all")
        database = self.database([("x|y", "1", "z"), ("x", "2", "y|z")], [("x|y", "a", "z"), ("x", "b", "y|z")], [("GENDER", "1", "a", "男"), ("GENDER", "2", "b", "女")])
        rows = self.execute(database, script)[0]
        self.assertEqual([row["check_status"] for row in rows], ["SAME", "SAME"])
        defaults["target_keys"] = ["region", "id"]
        swapped, _, _ = self.generate(defaults=defaults)
        database = self.database([("key", "1", "region")], [("region", "a", "key")], [("GENDER", "1", "a", "男")])
        self.assertEqual(self.execute(database, swapped)[0][0]["source_key_1"], "key")

    def test_file_mapping_mixed_domains_no_identity_assumption_and_duplicates(self):
        rows = [{"domain": "GENDER", "src": "1", "dst": "a", "label": "男"},
                {"domain": "ID", "src": 1, "dst": 100, "label": "身份证"}]
        path = self.file("mapping.json", rows)
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path), mapping_columns=self.mapping_columns)
        script, _, _ = self.generate()
        row = self.execute(self.database([("1", "1")], [("1", "a")]), script)[0][0]
        self.assertEqual(row["check_status"], "SAME")
        self.assertEqual(row["expected_target_code"], "a")
        rows.append(rows[0])
        path.write_text(json.dumps(rows), encoding="utf-8")
        self.assert_failure("DUPLICATE_MAPPING")

    def test_file_duplicate_validation_uses_active_conditions_not_unused_columns(self):
        rows = [{"domain": "GENDER", "src": "1", "dst": "a", "label": "男", "region": "NORTH", "enabled": "Y", "unused": "x"},
                {"domain": "GENDER", "src": "1", "dst": "b", "label": "女", "region": "NORTH", "enabled": "N", "unused": "y"},
                {"domain": "GENDER", "src": "1", "dst": "b", "label": "女", "region": "SOUTH", "enabled": "Y", "unused": "z"}]
        path = self.file("mapping.json", rows)
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path),
                           mapping_columns={**self.mapping_columns, "region": "region", "enabled": "enabled", "unused": "unused"},
                           defaults={**self.params["defaults"], "mapping_conditions": [
                               {"mapping_field": "region", "operator": "eq", "value_source": "source_field", "field": "region"},
                               {"mapping_field": "enabled", "operator": "eq", "value": "Y"}]})
        script, _, _ = self.generate()
        row = self.execute(self.database([("1", "1", "NORTH")], [("1", "a")]), script)[0][0]
        self.assertEqual(row["check_status"], "SAME")
        rows.append({**rows[0], "unused": "different"})
        path.write_text(json.dumps(rows), encoding="utf-8")
        self.assert_failure("DUPLICATE_MAPPING")

    def test_empty_file_domain_produces_typed_missing_relation(self):
        path = self.file("mapping.json", [{"domain": "DOCUMENT", "src": 1, "dst": 2, "label": "测试数据"}])
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path), mapping_columns=self.mapping_columns)
        script, _, _ = self.generate()
        row = self.execute(self.database([("1", "1")], [("1", "a")]), script)[0][0]
        self.assertEqual(row["check_status"], "MAPPING_MISSING")
        self.assertEqual(row["mapping_match_count"], 0)

    def test_file_text_domain_conversion_matches_the_plan(self):
        self.params.pop("code_table")
        path = self.file("mapping.json", [])
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path),
                           mapping_columns=self.mapping_columns, numeric_text_policy="convert")
        scenarios = [(1.0, 1.0, "SAME"), ("1.0", 1.0, "SAME"), (0.0, 0.0, "SAME"),
                     (-0.0, -0.0, "SAME"), (1.25, 1.25, "SAME"), (1, 1, "SAME"),
                     ("01", "01", "SAME"), ("1", 1.0, "MAPPING_MISSING"),
                     ("0.0", -0.0, "MAPPING_MISSING")]
        for plan_domain, mapping_domain, status in scenarios:
            with self.subTest(plan_domain=plan_domain, mapping_domain=mapping_domain):
                self.plan.write_text(json.dumps([{"src": "a", "src_field": "code", "dst": "b",
                                                 "dst_field": "code", "domain": plan_domain}]), encoding="utf-8")
                path.write_text(json.dumps([{"domain": mapping_domain, "src": "1", "dst": "a", "label": "男"},
                                            {"domain": "ID", "src": 1, "dst": 100, "label": "身份证"}]), encoding="utf-8")
                script, _, _ = self.generate()
                row = self.execute(self.database([("1", "1")], [("1", "a")]), script)[0][0]
                self.assertEqual(row["check_status"], status)

    def test_file_numeric_cell_in_text_domain_is_not_silently_discarded(self):
        self.plan.write_text(json.dumps([{"src": "a", "src_field": "code", "dst": "b",
                                         "dst_field": "code", "domain": "1.0"}]), encoding="utf-8")
        path = self.file("mapping.json", [{"domain": 1.0, "src": "1", "dst": "a", "label": "男"}])
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path),
                           mapping_columns=self.mapping_columns)
        self.assert_failure("INVALID_PARAMS")

    def test_file_numeric_domain_keeps_numeric_equality(self):
        self.params.pop("code_table")
        path = self.file("mapping.json", [])
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path),
                           mapping_columns=self.mapping_columns,
                           defaults={**self.params["defaults"], "types": {"domain": "numeric"}})
        for plan_domain, mapping_domain in (("01.0", 1.0), (0, "-0.0"), ("1.25", 1.25)):
            with self.subTest(plan_domain=plan_domain, mapping_domain=mapping_domain):
                self.plan.write_text(json.dumps([{"src": "a", "src_field": "code", "dst": "b",
                                                 "dst_field": "code", "domain": plan_domain}]), encoding="utf-8")
                path.write_text(json.dumps([{"domain": mapping_domain, "src": "1", "dst": "a", "label": "男"},
                                            {"domain": 100, "src": "1", "dst": "x", "label": "其他域"}]), encoding="utf-8")
                script, _, _ = self.generate()
                self.assertEqual(self.execute(self.database([("1", "1")], [("1", "a")]), script)[0][0]["check_status"], "SAME")

    def test_oracle_projected_columns_are_quoted_in_filters_and_summaries(self):
        configurations = [
            {"output_mode": "select_only", "result_scope": "errors_only", "include_summary": True},
            {"output_mode": "insert_only", "store_scope": "all", "include_summary": True},
            {"output_mode": "insert_only", "store_scope": "errors_only", "include_summary": True},
            {"output_mode": "select_and_insert", "store_scope": "all", "result_scope": "errors_only", "include_summary": True},
            {"output_mode": "select_and_insert", "store_scope": "errors_only", "result_scope": "errors_only"},
        ]
        for configuration in configurations:
            for layout in ("separate", "union_all"):
                with self.subTest(configuration=configuration, layout=layout):
                    script, _, _ = self.generate(sql_engine="oracle", engine_version="12.2",
                                                 statement_layout=layout, include_reverse_check=True, **configuration)
                    if configuration.get("result_scope") == "errors_only" or configuration.get("store_scope") == "errors_only":
                        self.assertIn('WHERE "check_status" <> \'SAME\'', script)
                    if configuration.get("include_summary"):
                        self.assertIn('SUM(CASE WHEN "check_status" = \'SAME\'', script)
                        self.assertIn('GROUP BY CAST("case_id" AS VARCHAR2(4000 CHAR)) COLLATE', script)
                        self.assertIn('c."case_id"', script)
                        self.assertIn('t."check_direction"', script)
                        self.assertIn('t."same_rows"', script)
                        self.assertNotIn("GROUP BY case_id", script)
                    self.assertIn("q.check_status", script)
                    self.assertIn('base."id" AS "key1"', script)
                    self.assertIn('CAST(q."key1" AS', script)
                    self.assertNotIn("q.key1", script)

    def test_internal_relation_names_do_not_shadow_input_tables(self):
        names = ["all_results", "cv1_base", "cv1_source", "cv1_mapping", "cv1_counted", "cv1_mapped",
                 "cv1_checked", "cv1_forward", "cv1_reverse_base", "cv1_reverse_checked", "cv1_reverse",
                 "summary_cases", "summary_counts", "saved_results", "CV1_BASE", "CVCHECK_CV1_BASE"]
        for name in names:
            for role in ("source", "target", "mapping"):
                with self.subTest(name=name, role=role):
                    source = name if role == "source" else "a"
                    target = name if role == "target" else "b"
                    self.plan.write_text(json.dumps([{"src": source, "src_field": "code", "dst": target,
                                                     "dst_field": "code", "domain": "GENDER"}]), encoding="utf-8")
                    code_table = {"table_parts": [name if role == "mapping" else "mapping"], "columns": self.mapping_columns}
                    script, _, _ = self.generate(code_table=code_table, include_summary=True, include_reverse_check=True)
                    database = self.database([("1", "1"), ("missing", "1")], [("1", "a")], [("GENDER", "1", "a", "男")])
                    original = {"source": "a", "target": "b", "mapping": "mapping"}[role]
                    database.execute('ALTER TABLE "' + original + '" RENAME TO "' + name + '"')
                    results = self.execute(database, script)
                    self.assertCountEqual([row["check_status"] for row in results[0]], ["SAME", "TARGET_MISSING"])
                    self.assertEqual(sum(row["total_rows"] for row in results[1]), 2)

    def test_internal_namespace_avoids_all_cases_result_and_inline_relations(self):
        self.plan.write_text(json.dumps([
            {"src": "cvcheck_all_results", "src_field": "code", "dst": "cvcheck1_all_results", "dst_field": "code", "domain": "GENDER", "case_id": "first"},
            {"src": "cvcheck_all_results", "src_field": "code", "dst": "cvcheck1_all_results", "dst_field": "code", "domain": "GENDER", "case_id": "second"},
        ]), encoding="utf-8")
        self.params["test_plan_columns"]["case_id"] = "case_id"
        path = self.file("mapping.json", [{"domain": "GENDER", "src": "1", "dst": "a", "label": "男"}])
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path), mapping_columns=self.mapping_columns)
        for mode in ("select_only", "insert_only", "select_and_insert"):
            for layout in ("separate", "union_all"):
                for result_table in ("cv2_mapping", "saved_results", "summary_counts", "cvcheck2_all_results"):
                    with self.subTest(mode=mode, layout=layout, result_table=result_table):
                        changes = {"result_scope" if mode == "select_only" else "store_scope": "errors_only"}
                        if mode != "select_only":
                            changes["result_table"] = result_table
                        if mode == "select_and_insert":
                            changes.update(store_scope="all", result_scope="errors_only")
                        script, plan, _ = self.generate(output_mode=mode, statement_layout=layout,
                                                        include_summary=True, include_reverse_check=True, **changes)
                        database = self.database([("ok", "1"), ("bad", "1"), ("missing", "1")], [("ok", "a"), ("bad", "b")])
                        database.executescript('ALTER TABLE a RENAME TO cvcheck_all_results; ALTER TABLE b RENAME TO cvcheck1_all_results;')
                        results = self.execute(database, script)
                        details = ([dict(row) for row in database.execute('SELECT * FROM "' + plan["result_table"][0] + '"')]
                                   if mode == "insert_only" else [row for group in results for row in group if "check_status" in row])
                        self.assertCountEqual([row["check_status"] for row in details], ["DIFFERENT", "TARGET_MISSING"] * 2)
                        summaries = [row for group in results for row in group if "total_rows" in row]
                        self.assertEqual(sum(row["total_rows"] for row in summaries), 6)

    def test_dialect_identifier_length_boundaries(self):
        boundaries = [("postgresql", "12", "x" * 63, "x" * 64),
                      ("postgresql", "12", "码" * 21, "码" * 21 + "x"),
                      ("mysql", "8", "码" * 64, "码" * 65),
                      ("mariadb", "10.2", "x" * 64, "x" * 65),
                      ("oracle", "12.2", "x" * 128, "码" * 43),
                      ("oracle", "12.2", "码" * 42 + "xx", "码" * 42 + "xxx"),
                      ("sqlserver", "14", "码" * 128, "码" * 129),
                      ("sqlserver", "14", "😀" * 64, "😀" * 65),
                      ("sqlite", "3.35", "码" * 128, "码" * 129),
                      ("trino", "483", "码" * 128, "码" * 129)]
        for engine, version, allowed, rejected in boundaries:
            with self.subTest(engine=engine, allowed=allowed, rejected=rejected):
                dialect = HELPERS.Dialect(engine, version)
                self.assertIn(allowed, dialect.quote(allowed))
                self.assertIn(allowed, dialect.table(["schema", allowed]))
                with self.assertRaises(PluginError):
                    dialect.quote(rejected)
                with self.assertRaises(PluginError):
                    dialect.table([rejected, "table"])

    def test_sql_identifiers_rejected_before_artifacts_are_written(self):
        long_name = "x" * 64
        original = json.loads(self.plan.read_text())
        for field in ("src", "src_field", "dst", "dst_field"):
            with self.subTest(field=field):
                self.plan.write_text(json.dumps([{**original[0], field: long_name}]), encoding="utf-8")
                self.assert_failure("INVALID_PARAMS", sql_engine="postgresql", engine_version="12")
        self.plan.write_text(json.dumps(original), encoding="utf-8")
        configurations = [
            {"code_table": {"table_parts": [long_name, "mapping"], "columns": self.mapping_columns}},
            {"code_table": {"table_parts": ["mapping"], "columns": {**self.mapping_columns, "unused": long_name}}},
            {"defaults": {"source_keys": [long_name], "target_keys": ["id"]}},
            {"defaults": {"source_keys": ["id"], "target_keys": [long_name]}},
            {"defaults": {**self.params["defaults"], "source_filters": [{"field": long_name, "operator": "is_null"}]}},
            {"defaults": {**self.params["defaults"], "mapping_conditions": [{"mapping_field": "region", "operator": "eq", "value_source": "source_field", "field": long_name}]},
             "code_table": {"table_parts": ["mapping"], "columns": {**self.mapping_columns, "region": "region"}}},
            {"output_mode": "insert_only", "result_table": long_name},
            {"output_mode": "insert_only", "result_columns": {"source_code": long_name}},
            {"output_mode": "insert_only", "result_columns": {"source_code": "audit_" + "x" * 57 + "A", "actual_target_code": "audit_" + "x" * 57 + "B"}},
        ]
        for configuration in configurations:
            with self.subTest(configuration=configuration):
                self.assert_failure("INVALID_PARAMS", sql_engine="postgresql", engine_version="12", **configuration)
        self.assert_failure("INVALID_PARAMS", sql_engine="mysql", engine_version="8", output_mode="insert_only", result_table="x" * 65)

    def test_trino_mysql_sqlserver_oracle_boundaries_and_escaping(self):
        for engine, version in (("mysql", "8"), ("mariadb", "10.2"), ("postgresql", "12"), ("sqlserver", "14"), ("oracle", "12.2")):
            for mode in ("insert_only", "select_and_insert"):
                script, _, _ = self.generate(sql_engine=engine, engine_version=version, output_mode=mode)
                self.assertTrue(script.startswith("CREATE TABLE"))
                self.assertIn("INSERT INTO", script)
        self.assert_failure("UNSUPPORTED_COMBINATION", sql_engine="oracle", engine_version="12.2", output_mode="insert_only", create_policy="create_if_missing")
        postgres = HELPERS.Dialect("postgresql", "12")
        self.assertEqual(postgres.literal("a\\b'c"), "E'a\\\\b''c'")
        self.assertIn("DATALENGTH", HELPERS.Dialect("sqlserver", "14").equal("left_value", "right_value", "text"))
        for engine, version in (("mysql", "8"), ("mariadb", "10.2")):
            dialect = HELPERS.Dialect(engine, version)
            self.assertIn("AS SIGNED", dialect.typed("candidate_count", "count"))
            self.assertIn("CHAR CHARACTER SET utf8mb4", dialect.text("source_code"))
        self.assertIn("AS TEXT", HELPERS.Dialect("sqlite", "3.35").inline(None, "text"))
        for value in ("1e999999999", "123456789012345678901", "0.1234567890123456789"):
            with self.subTest(value=value), self.assertRaises(PluginError):
                HELPERS.value_type(value, "numeric", {})

    def test_same_file_rejected_before_two_snapshots_are_created(self):
        params = {**self.params, "mapping_mode": "file", "mapping_source": "separate_file", "mapping_input": str(self.plan)}
        params.pop("code_table")
        with self.assertRaisesRegex(ValueError, "重复来源文件"):
            self.runtime.run("sql.codecheck", params)

    def test_heterogeneous_keys_and_dimensions_share_union_projection(self):
        rows = json.loads(self.plan.read_text())
        rows = [{**rows[0], "case_id": "one", "src_keys": ["id"], "dst_keys": ["id"]},
                {**rows[0], "case_id": "two", "src_keys": ["id", "region"], "dst_keys": ["id", "region"]}]
        self.plan.write_text(json.dumps(rows), encoding="utf-8")
        columns = {**self.plan_columns, "case_id": "case_id", "source_keys": "src_keys", "target_keys": "dst_keys"}
        script, _, _ = self.generate(test_plan_columns=columns, statement_layout="union_all", include_summary=True)
        rows = self.execute(self.database([("1", "1", "NORTH")], [("1", "a", "NORTH")], [("GENDER", "1", "a", "男")]), script)[0]
        self.assertEqual(len(rows), 2)
        self.assertIsNone(rows[0]["source_key_2"])
        self.assertEqual(rows[1]["source_key_2"], "NORTH")
        self.assertTrue(all(row["check_status"] == "SAME" for row in rows))

    def test_case_sensitive_matching_without_trim_or_code_fallback(self):
        script, _, _ = self.generate()
        for source, expected in (("01", "MAPPING_MISSING"), ("1 ", "MAPPING_MISSING"), ("1", "DIFFERENT")):
            row = self.execute(self.database([("key", source)], [("key", "A")], [("GENDER", "1", "a", "男")]), script)[0][0]
            self.assertEqual(row["check_status"], expected)
        row = self.execute(self.database([("KEY", "1")], [("key", "a")], [("GENDER", "1", "a", "男")]), script)[0][0]
        self.assertEqual(row["check_status"], "SOURCE_MISSING")

    def test_result_columns_and_business_table_write_protection(self):
        script, _, _ = self.generate(output_mode="select_and_insert", result_table="audit_results", result_columns={"source_code": "源码值", "actual_target_code": "实际码值"})
        database = self.database([("1", "1")], [("1", "a")], [("GENDER", "1", "a", "男")])
        self.assertEqual(self.execute(database, script)[0][0]["source_code"], "1")
        self.assertEqual(database.execute('SELECT "源码值" FROM audit_results').fetchone()[0], "1")
        self.assertEqual(database.execute("SELECT count(*) FROM a").fetchone()[0], 1)
        for name in ("a", "b", "mapping", "main.a"):
            with self.subTest(name=name):
                self.assert_failure("INVALID_PARAMS", output_mode="insert_only", result_table=name)

    def assert_failure(self, code, **changes):
        params = copy.deepcopy(self.params)
        params.update(changes)
        task, result = self.runtime.run("sql.codecheck", params)
        self.assertEqual(result.status, "failed", result.to_dict())
        self.assertEqual(result.data.get("error_code"), code, result.to_dict())
        self.assertEqual(result.files, [])
        self.assertFalse((Path(self.runtime.get_task(task)["workspace_path"]) / "output/codecheck.sql").exists())

    def test_configuration_rejects_ambiguous_or_unsafe_combinations(self):
        for changes in ({"store_scope": "all"}, {"output_mode": "insert_only", "result_scope": "all"},
                        {"output_mode": "select_and_insert", "store_scope": "errors_only"},
                        {"union_batch_size": 2}, {"mapping_source": "same_workbook"},
                        {"defaults": {"source_keys": [], "target_keys": ["id"]}},
                        {"defaults": {"source_keys": ["id"], "target_keys": ["id", "code"]}},
                        {"code_table": {"table_parts": ["mapping;DROP"], "columns": self.mapping_columns}},
                        {"defaults": {**self.params["defaults"], "mapping_conditions": [{"mapping_field": "region", "operator": "eq", "value": None}]}},
                        {"defaults": {**self.params["defaults"], "mapping_conditions": [{"mapping_field": "region", "operator": "eq", "field": "region", "value": "x"}]}},
                        {"defaults": {**self.params["defaults"], "mapping_conditions": [{"mapping_field": "region", "operator": "is_null", "value": "x"}]}}):
            with self.subTest(changes=changes):
                self.assert_failure("INVALID_PARAMS", **changes)

    def test_limits_and_dialect_rendering_are_not_execution_claims(self):
        for engine, version, quote in (("postgresql", "12", '"a"'), ("mysql", "8", "`a`"), ("mariadb", "10.2", "`a`"), ("oracle", "12.2", '"a"'), ("sqlserver", "14", "[a]"), ("trino", "483", '"a"')):
            with self.subTest(engine=engine):
                script, _, result = self.generate(sql_engine=engine, engine_version=version)
                self.assertIn(quote, script)
                self.assertIn("COUNT", script)
                self.assertTrue(any("未在实际数据库" in warning for warning in result.warnings))
        self.assert_failure("UNSUPPORTED_COMBINATION", sql_engine="trino", engine_version="483", output_mode="insert_only")
        self.assert_failure("UNSUPPORTED_ENGINE", engine_version="3.1")
        self.assert_failure("OUTPUT_LIMIT_EXCEEDED", max_sql_bytes=1024)
        with self.assertRaises(PluginError):
            HELPERS.Dialect("elasticsearch", "8")

    def test_duplicate_cases_merge_and_conflicting_ids(self):
        original = json.loads(self.plan.read_text())
        self.plan.write_text(json.dumps(original * 3), encoding="utf-8")
        self.assert_failure("DUPLICATE_CASE_CONFIG")
        _, plan, result = self.generate(duplicate_case_policy="merge_identical")
        self.assertEqual(result.data["merged_rows"], 2)
        self.assertEqual(len(plan["cases"][0]["locations"]), 3)
        rows = [{**original[0], "case_id": case_id} for case_id in ("one", "two", "three")]
        self.plan.write_text(json.dumps(rows), encoding="utf-8")
        self.params["test_plan_columns"]["case_id"] = "case_id"
        _, plan, _ = self.generate()
        self.assertEqual(len(plan["cases"]), 3)

    def test_file_null_token_and_numeric_declarations(self):
        path = self.file("mapping.json", [{"domain": "GENDER", "src": "1", "dst": "__NULL__", "label": None}])
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_source="separate_file", mapping_input=str(path), mapping_columns=self.mapping_columns)
        script, _, _ = self.generate(allow_mapping_null=True, null_token="__NULL__")
        row = self.execute(self.database([("1", "1")], [("1", None)]), script)[0][0]
        self.assertEqual(row["check_status"], "NULL_VALUE")
        path.write_text(json.dumps([{"domain": "GENDER", "src": 1, "dst": 2, "label": None}]), encoding="utf-8")
        self.assert_failure("INVALID_PARAMS")
        script, _, _ = self.generate(numeric_text_policy="convert")
        self.assertEqual(self.execute(self.database([("1", "1")], [("1", "2")]), script)[0][0]["check_status"], "SAME")
        defaults = {**self.params["defaults"], "types": {"source_code": "numeric", "target_code": "numeric"}}
        script, _, _ = self.generate(defaults=defaults)
        database = self.database()
        database.executescript("DROP TABLE a; DROP TABLE b; CREATE TABLE a(id TEXT, code NUMERIC); CREATE TABLE b(id TEXT, code NUMERIC); INSERT INTO a VALUES ('1',1); INSERT INTO b VALUES ('1',2);")
        self.assertEqual(self.execute(database, script)[0][0]["check_status"], "SAME")

    def test_excel_layouts_equivalent_and_originals_unchanged(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("Excel optional dependency not installed")
        self.params.pop("code_table")
        self.params.update(mapping_mode="file", mapping_columns=self.mapping_columns)
        plans = [["src", "src_field", "dst", "dst_field", "domain"], ["a", "code", "b", "code", "GENDER"]]
        mappings = [["domain", "src", "dst", "label"], ["GENDER", "1", "a", "男"], ["GENDER", "2", "b", "女"]]
        for layout in ("separate_files", "two_sheets", "side_blocks", "stacked_blocks", "combined_rows"):
            with self.subTest(layout=layout):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "清单"
                options, mapping_options = {"sheet": "清单"}, {"sheet": "清单"}
                if layout == "combined_rows":
                    sheet.append(plans[0] + ["map_src", "map_dst", "label"])
                    for row in mappings[1:]:
                        sheet.append(plans[1] + row[1:])
                    columns = {**self.mapping_columns, "source_code": "map_src", "target_code": "map_dst"}
                else:
                    for row in plans:
                        sheet.append(row)
                    columns = self.mapping_columns
                    if layout == "two_sheets":
                        mapped_sheet = workbook.create_sheet("映射")
                        for row in mappings:
                            mapped_sheet.append(row)
                        mapping_options = {"sheet": "映射"}
                    elif layout == "side_blocks":
                        for row_number, row in enumerate(mappings, 1):
                            for column_number, value in enumerate(row, 7):
                                sheet.cell(row_number, column_number, value)
                        options["excel_range"] = "A1:E2"
                        mapping_options["excel_range"] = "G1:J3"
                    elif layout == "stacked_blocks":
                        for row_number, row in enumerate(mappings, 5):
                            for column_number, value in enumerate(row, 1):
                                sheet.cell(row_number, column_number, value)
                        options["excel_range"] = "A1:E2"
                        mapping_options["excel_range"] = "A5:D7"
                path = self.root / (layout + ".xlsx")
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                changes = {"test_plan_input": str(path), "test_plan_options": options, "mapping_options": mapping_options,
                           "mapping_source": "same_workbook", "mapping_columns": columns}
                if layout == "combined_rows":
                    changes["file_layout"] = "combined_rows"
                if layout == "separate_files":
                    map_book = Workbook()
                    for row in mappings:
                        map_book.active.append(row)
                    map_book.active.title = "清单"
                    map_path = self.root / "separate_mapping.xlsx"
                    map_book.save(map_path)
                    map_book.close()
                    map_before = map_path.read_bytes()
                    changes.update(mapping_source="separate_file", mapping_input=str(map_path))
                script, plan, _ = self.generate(**changes)
                rows = self.execute(self.database([("1", "1"), ("2", "2")], [("1", "a"), ("2", "b")]), script)[0]
                self.assertEqual([row["check_status"] for row in rows], ["SAME", "SAME"])
                self.assertEqual(len(plan["cases"]), 1)
                self.assertEqual(path.read_bytes(), before)
                if layout == "separate_files":
                    self.assertEqual(map_path.read_bytes(), map_before)

    def test_excel_region_is_bounded_and_has_absolute_coordinates(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("Excel optional dependency not installed")
        book = Workbook()
        book.active["C4"] = "id"
        book.active["D4"] = "value"
        book.active["C5"] = "001"
        book.active["D5"] = "a"
        book.active["Z99"] = "outside"
        path = self.root / "region.xlsx"
        book.save(path)
        book.close()
        dataset = read_dataset(path, {"excel_range": "C4:D5"})
        self.assertEqual(dataset["rows"], [{"id": "001", "value": "a"}])
        self.assertEqual(dataset["locations"][0]["row"], 5)
        self.assertEqual(dataset["locations"][0]["first_column"], 3)
        for options in ({"excel_range": "C:D"}, {"excel_range": "C4:D5", "header_row": 1},
                        {"excel_range": "C4:D5", "start_row": 7}, {"excel_range": "A1:XFE2"},
                        {"excel_range": "A1:D50", "max_rows": 2}):
            with self.subTest(options=options), self.assertRaises(PluginError):
                read_dataset(path, options)
        with self.assertRaises(PluginError):
            read_dataset(self.plan, {"excel_range": "A1:B2"})


if __name__ == "__main__":
    unittest.main()
