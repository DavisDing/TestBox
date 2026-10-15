"""Runtime/Host regressions for the 2026-10-15 code audit."""
from __future__ import annotations

import json
from contextlib import closing
from pathlib import Path
import shutil
import sqlite3
import tempfile
import unittest

from testbox.core.runtime import Runtime

ROOT = Path(__file__).resolve().parents[1]


class AuditRuntimeTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="testbox-audit-fix-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        shutil.copytree(ROOT / "plugins", self.root / "plugins",
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
        self.runtime = Runtime(self.root)
        self.addCleanup(self.runtime.close)

    def run_success(self, command, params):
        task, result = self.runtime.run(command, params)
        self.assertEqual(result.status, "success", result.message)
        return task, result

    def output(self, task, result):
        return self.runtime.get_task_output_path(task, result.files[0])

    def test_sql_headers_inside_literals_and_quoted_identifiers_are_ignored(self):
        source = self.root / "schema.sql"
        source.write_text('''CREATE TABLE "real_table" (
            id INT,
            note VARCHAR(150) DEFAULT 'prefix ''CREATE TABLE fake_table (x INT)'' suffix',
            description VARCHAR(150) COMMENT 'CREATE TABLE comment_table (x INT)'
        );
        CREATE TABLE `second_table` (`CREATE TABLE misleading` INT);
        CREATE TABLE [third_table] (id INT);''', encoding="utf-8")
        for strict in (False, True):
            with self.subTest(strict=strict):
                task, result = self.run_success("sql.parse", {
                    "input": str(source), "format": "json", "fail_on_unsupported": strict,
                })
                rows = json.loads(self.output(task, result).read_text(encoding="utf-8"))
                self.assertEqual([(row["table"], row["field"]) for row in rows], [
                    ("real_table", "id"), ("real_table", "note"), ("real_table", "description"),
                    ("second_table", "CREATE TABLE misleading"), ("third_table", "id"),
                ])

    def test_sparse_excel_import_preserves_constraints_and_empty_comments(self):
        try:
            from openpyxl import Workbook
        except ModuleNotFoundError:
            self.skipTest("openpyxl is not installed")
        source = self.root / "fields.xlsx"
        workbook = Workbook()
        workbook.active.append(["table", "field", "type", "comment", "primary_key", "nullable", "default"])
        workbook.active.append(["sample", "id", "INTEGER", None, True, False, "99"])
        workbook.active.append(["sample", "note", "VARCHAR(8)", None, None, True, None])
        workbook.save(source)
        workbook.close()
        _, result = self.run_success("data.mock", {
            "count": 1, "format": "json", "source_file": str(source),
            "source_format": "excel", "preview": True,
        })
        first, second = result.data["tables"][0]["fields"]
        self.assertTrue(first["primary_key"])
        self.assertFalse(first["nullable"])
        self.assertEqual(first["default"], "99")
        self.assertFalse(first.get("comment"))
        self.assertEqual(second["name"], "note")
        self.assertFalse(second.get("primary_key", False))
        self.assertTrue(second["nullable"])

    def test_imported_composite_primary_key_script_executes_in_sqlite(self):
        source = self.root / "pair.sql"
        source.write_text("CREATE TABLE pair (a INTEGER, b INTEGER, PRIMARY KEY(a,b));", encoding="utf-8")
        task, result = self.run_success("data.mock", {
            "count": 2, "seed": 7, "format": "sql", "source_file": str(source),
            "source_format": "sql", "sql_dialect": "sqlite", "sql_create_table": True,
        })
        script = self.output(task, result).read_text(encoding="utf-8")
        with closing(sqlite3.connect(":memory:")) as database:
            database.executescript(script)
            self.assertEqual([row[5] for row in database.execute("PRAGMA table_info(pair)")], [1, 2])
            # Members of a composite primary key are not individually UNIQUE.
            database.execute("INSERT INTO pair VALUES (1000001,1000002),(1000001,1000003),(1000004,1000002)")
            with self.assertRaises(sqlite3.IntegrityError):
                database.execute("INSERT INTO pair VALUES (1000001,1000002)")
        self.assertEqual(script.count("PRIMARY KEY"), 1)

    def test_dotted_column_parse_select_roundtrip_executes_with_qualified_table(self):
        source = self.root / "dotted.sql"
        ddl = 'CREATE TABLE main.sample ("a.b" INTEGER);'
        source.write_text(ddl, encoding="utf-8")
        parsed_task, parsed = self.run_success("sql.parse", {"input": str(source), "format": "json"})
        task, result = self.run_success("sql.select", {
            "input": str(self.output(parsed_task, parsed)), "dialect": "sqlite",
        })
        statement = self.output(task, result).read_text(encoding="utf-8")
        with closing(sqlite3.connect(":memory:")) as database:
            database.executescript(ddl + 'INSERT INTO sample ("a.b") VALUES (42);')
            self.assertEqual(database.execute(statement).fetchall(), [(42,)])
        self.assertIn('FROM "main"."sample"', statement)
        _, unsupported = self.runtime.run("sql.select", {
            "input": str(self.output(parsed_task, parsed)), "dialect": "auto",
        })
        self.assertEqual(unsupported.status, "failed")
        self.assertIn("auto", unsupported.message)

    def test_address_filters_choose_only_matching_provinces_and_cities(self):
        cases = [
            ({"city": ["深圳市"]}, ["广东省深圳市"]),
            ({"district": ["南山区"]}, ["广东省深圳市南山区", "黑龙江省鹤岗市南山区"]),
            ({"province": ["广东省", "浙江省"], "city": ["深圳市"]}, ["广东省深圳市"]),
            ({"city": ["深圳市", "杭州市"]}, ["广东省深圳市", "浙江省杭州市"]),
        ]
        for options, prefixes in cases:
            with self.subTest(options=options):
                params = {"count": 20, "seed": 1, "format": "json", "fields": [{
                    "name": "address", "generator": "china_address",
                    "options": {**options, "street_mode": "none"},
                }]}
                first_task, first = self.run_success("data.mock", params)
                second_task, second = self.run_success("data.mock", params)
                content = self.output(first_task, first).read_bytes()
                self.assertEqual(content, self.output(second_task, second).read_bytes())
                self.assertTrue(all(any(row["address"].startswith(prefix) for prefix in prefixes)
                                    for row in json.loads(content)))

    def test_address_contradictory_filters_fail_explicitly(self):
        _, result = self.runtime.run("data.mock", {"count": 1, "format": "json", "fields": [{
            "name": "address", "generator": "china_address",
            "options": {"province": ["浙江省"], "city": ["深圳市"]},
        }]})
        self.assertEqual(result.status, "failed")
        self.assertIn("不匹配", result.message)

    def test_evidence_blank_rows_do_not_consume_screenshots_or_excel_status(self):
        try:
            from openpyxl import Workbook, load_workbook
            from openpyxl.styles import PatternFill
            from PIL import Image
            import docx  # noqa: F401: report generation dependency
        except ModuleNotFoundError:
            self.skipTest("evidence dependencies are not installed")
        source = self.root / "cases.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["测试名称", "验证点", "步骤名称", "测试结果"])
        sheet.append(["One", "check", "first", None])
        sheet.append([None, None, None, None])
        sheet.append([None, None, "continued", None])
        sheet.append([None, None, None, None])
        sheet.append(["Two", "other", None, None])
        sheet["A8"].fill = PatternFill("solid", fgColor="EEEEEE")
        workbook.save(source)
        workbook.close()
        screenshots = []
        for index in range(3):
            image = self.root / f"shot-{index}.png"
            Image.new("RGB", (100, 60), "white").save(image)
            screenshots.append(str(image))
        task, result = self.run_success("evidence.build", {
            "input": str(source), "interactive": False, "screenshots": screenshots,
        })
        self.assertEqual(result.data["rows"], [2, 4, 6])
        self.assertEqual([item["case_name"] for item in result.data["items"]], ["One", "One", "Two"])
        self.assertEqual([item["screenshot"] for item in result.data["items"]], [f"shot-{i}.png" for i in range(3)])
        self.assertEqual(result.data["pending_count"], 0)
        updated = load_workbook(self.runtime.get_task_output_path(task, "executed-cases.xlsx"))
        self.addCleanup(updated.close)
        for row in (2, 4, 6):
            self.assertEqual(updated.active.cell(row, 4).value, "已执行")
        for row in (3, 5, 7, 8):
            self.assertIsNone(updated.active.cell(row, 4).value)


if __name__ == "__main__":
    unittest.main()
